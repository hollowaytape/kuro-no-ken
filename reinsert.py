"""
    Kuro no Ken reinserter.
"""

import os
import sys
from shutil import copyfile

# bsd_tool, fix_pointers and script_map live in tools/
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'tools'))

from romtools.disk import Disk, Gamefile, Block
from romtools.dump import PointerExcel
from ordered_dump import OrderedDumpExcel
from pointer_edit import KuroGamefile
from rominfo import SRC_DISK, DEST_DISK, FILES, FILES_TO_REINSERT, COMPRESSED_FILES_TO_EDIT, ARCHIVES_TO_REINSERT
from rominfo import FILE_BLOCKS, LENGTH_SENSITIVE_BLOCKS, DECOMPRESSED_SIZE_LIMITS, NAMES
from rominfo import  DUMP_XLS_PATH, POINTER_XLS_PATH, POINTERS_TO_REASSIGN, CONTROL_CODES
from rominfo import MAPPING_BUILD
from rominfo import BSD_FILES_WITH_TEXT, BSD_DECOMPRESSED_SIZE_LIMIT

from asm import BYTE_EDITS
from fa1 import repack, unpack
from bsd_tool import parse_bsd, reinsert_bsd
from decompress import decompress
import fix_pointers
import script_decode

Dump = OrderedDumpExcel(DUMP_XLS_PATH)   # row order in the sheet is free
PtrDump = PointerExcel(POINTER_XLS_PATH)

# Area hub scripts whose only text is the centered location name. Their pointer
# tables can't be trusted (the regex finder misreads the hub's fixed-size script
# units), so their strings must keep the exact Japanese byte length: they skip
# typeset() (which collapses the centering spaces) and get space-padded.
FIXED_LENGTH_SCN = {'02OLB.SCN', '03YSK.SCN'}

# Map scripts with no `09 <addr>` entry table (12MRS.SCN, 01FLD.SCN...) can't be decoded
# reliably, so gen_pointers gives them no pointer sheet - but they end in a zone table of
# block addresses, and any length change strands it (the crash #4 class). Until their
# zone tables are located, their strings keep the exact Japanese length too; an English
# line that doesn't fit stops the build with a message instead of crashing the game.
def _no_entry_table(fn):
    from script_map import slot_base
    path = os.path.join('original', 'decompressed', fn)
    return os.path.exists(path) and slot_base(open(path, 'rb').read())[0] is None

# 00IPL.SCN and 02OLB00A.SCN have no entry table either, but they are hand-listed scripts
# whose pointers were curated and verified in game, so the rule is only for the rest.
# (07CSLI0x, 09HIKI01, 25TOU00A/C are stranger still: they start mid-structure and end in
# x86 code, and no pointer in them decodes.) If one of these has a line that doesn't fit,
# the file is left in Japanese with a warning rather than stopping the whole build.
from rominfo import CURATED_FILES
UNKNOWN_FORMAT_SCN = {f for f in FILES_TO_REINSERT if f.endswith('.SCN') and _no_entry_table(f)
                      and f not in CURATED_FILES}
FIXED_LENGTH_SCN |= UNKNOWN_FORMAT_SCN
SKIPPED_TOO_LONG = []


_NO_TEXT = {}


def file_shows_no_text(filename):
    """True when not one of a file's dump rows sits inside a `40 02 <string> 00` print.

    A row outside every print can still be real - menu choices are drawn another way - but
    a file where *no* row is inside one displays nothing at all. 31END.SCN is the only one:
    the ending is x86 code (`e8 22 00` call, `eb 05` jmp), and the four strings the dumper
    found in it are stray bytes, one of them `ab ab ab ab` filler.
    """
    if filename not in _NO_TEXT:
        from script_map import strings
        path = os.path.join('original', 'decompressed', filename)
        try:
            said = strings(open(path, 'rb').read())
        except OSError:
            said = []
        rows = [t.location for t in translations_of(filename, include_blank=True,
                                                          sheet_name='SCNs')]
        _NO_TEXT[filename] = bool(rows) and bool(said) and not any(
            any(a - 4 <= o <= b for a, b in said) for o in rows)
    return _NO_TEXT[filename]


NAME_FILLED = []            # (filename, offset, english) for the report at the end
_NAME_MAP = {}



FORMULA_CELLS = []          # (filename, offset) for the report at the end


def translations_of(block, **kw):
    """Dump.get_translations, with any formula in the English column resolved.

    The translator's sheet holds formulas there on purpose: a name plate reads its English
    out of Names + Places, so a character's name is written down once. romtools loads the
    workbook without `data_only`, so what comes back for those cells is the formula itself
    - `='Names + Places'!C12` - and writing that into the game is exactly the failure
    tools/merge_dump.py describes for `=E14` references.

    The formula is resolved here rather than left to the blank-cell path further down,
    because everything in between measures the English: the check that decides whether a
    script can keep the Japanese's byte length runs on these very rows, and a formula read
    as an empty cell sailed past it and then failed hard when the fill put the name back
    ('Shinobu' is 7 bytes where 07CSLI01 has 6). Resolved here, the row carries exactly
    what the literal used to, and that script gets the same warning it always did.

    A formula that names nothing in the glossary becomes an empty cell, which falls back to
    the Japanese, as an untouched row does - never to the text of the formula itself.
    """
    out = Dump.get_translations(block, **kw)
    for t in out:
        if t.en_bytestring.startswith(b'='):
            FORMULA_CELLS.append((getattr(block, 'filename', None), t.location))
            t.en_bytestring = name_plate_english(t.jp_bytestring) or b''
    return out


def name_plate_english(japanese):
    """English for a cell that is *only* a character's name, or None.

    A name plate is the same string hundreds of times over - 171 rows already carry one,
    and 224 identical rows are blank - so it should be written down once and filled in on
    the way past. Two sources, glossary first:

    * **Names + Places** (rominfo.NAMES), which is where a name is decided. Its Japanese is
      the full name, so the part before the first `・` is indexed too: the plates say
      シノブ and カイエス where the sheet says シノブ・リュード and カイエス・ナインターク.
    * **the translator's own rows**: a short, unpunctuated Japanese cell that always has
      the same English wherever it is filled in (リーエ -> Lilie, 170 times) is that name,
      spelled the way they chose. Any Japanese with two different Englishes is left alone.

    Only a cell that is the whole name matches; a name inside a sentence stays the
    translator's to write.
    """
    if not _NAME_MAP:
        _NAME_MAP.update(_build_name_map())
    try:
        text = bytes(japanese).decode('cp932')
    except UnicodeDecodeError:
        return None
    return _NAME_MAP.get(text.strip().strip('\u3000').strip())


def _looks_like_plate(jp, en):
    """Short, unpunctuated, and the English is short too - a label, not a line."""
    return (jp and en and len(jp) <= 8 and len(en) <= 24
            and not any(c in jp for c in '\u300c\u300d\u3002\u3001\uff1f\uff01\u30fb\\')
            and not any(c in en for c in '.!?"'))


def _build_name_map():
    """-> {japanese name plate: english}, glossary over learned, conflicts dropped."""
    import collections
    learned = collections.defaultdict(set)
    try:
        import openpyxl
        from rominfo import sheet_columns
        wb = openpyxl.load_workbook(DUMP_XLS_PATH, read_only=True, data_only=True)
        ws = wb['SCNs']
        col = sheet_columns(ws)
        ji, ei = col.get('Japanese'), col.get('English')
        for row in ws.iter_rows(min_row=2, values_only=True):
            if ji is None or ei is None or len(row) <= max(ji, ei):
                continue
            jp = str(row[ji]).strip() if isinstance(row[ji], str) else ''
            en = str(row[ei]).strip() if isinstance(row[ei], str) else ''
            if _looks_like_plate(jp, en):
                learned[jp].add(en)
    except Exception:
        learned = {}
    out = {jp: next(iter(v)) for jp, v in learned.items() if len(v) == 1}
    for jp_full, en in getattr(NAMES, 'pairs', lambda: [])():
        out[jp_full] = en
        head = jp_full.split('\u30fb')[0]
        if head:
            out.setdefault(head, en)
    return {jp: en.encode('cp932', 'replace') for jp, en in out.items()}

def too_long_for_fixed(gf):
    """-> [(offset, english, jp_len)] for lines that can't keep the Japanese length."""
    probe = Block(gf, (0, len(gf.original_filestring)))
    return [(t.location, t.en_bytestring, len(t.jp_bytestring))
            for t in translations_of(probe, include_blank=True, sheet_name='SCNs')
            if t.en_bytestring and len(t.en_bytestring) > len(t.jp_bytestring)]

LINE_MAX = 48
INDENT = b'  '
NEWLINE_SEQ = b'\\n\x00\x40\x02'   # \n + rendering sync (5C 6E 00 40 02)
SPLIT_SEQ = b'\\f\x00;@\x02'       # \f page break (5C 66 00 3B 40 02)

# --gfx-dialogue: text inside a box frame is drawn with the variable-width font, so it is
# wrapped by pixel width (tools/vwf_metrics.py) instead of by characters. The box is 50
# columns = 400 px; like LINE_MAX (50 - 2), leave the 16 px of a speaker line's \i2.
VWF = '--gfx-dialogue' in sys.argv
VWF_LINE_PX = 50 * 8 - 16
TYPESET_COUNTS = {'vwf': 0, 'text': 0, 'unknown': 0}


def wrap_line(text, max_width=LINE_MAX, indent=INDENT, measure=len):
    """Word-wrap a single line of text, returning a list of wrapped lines.
    `measure` gives a line's width in the units of `max_width` (characters by default)."""
    if not text:
        return [indent]

    words = text.split(b' ')
    lines = []
    current = indent

    for word in words:
        if not word:
            continue
        needed = measure(current + (b' ' if current != indent else b'') + word)
        if needed > max_width and current != indent:
            lines.append(current.rstrip())
            current = indent
        if current == indent:
            current += word
        else:
            current += b' ' + word

    if current and current != indent:
        lines.append(current.rstrip())
    elif not lines:
        lines = [indent]

    return lines

BOX_ROWS = 5   # a text box shows a speaker name + 4 lines (02OLB00A 0x3bb's 6th row never showed)


def typeset(s, rows_before=0, vwf=False):
    """Format English text for the game's text box.

    Word-wraps to LINE_MAX characters per line (vwf: VWF_LINE_PX pixels) with INDENT prefix, honours
    [SPLIT] page breaks and existing newlines, and starts a new page ([SPLIT])
    whenever the box would run past BOX_ROWS rows. `rows_before` is how many
    rows of the current page earlier cells already used (tracked in the reinsert loop).
    Character names from NAMES are returned unchanged.

    Returns (bytes, rows used on the page the text ends on).
    """
    if s in NAMES:
        return s, rows_before + 1

    out_pages = []
    rows = rows_before
    for page_no, page in enumerate(s.split(SPLIT_SEQ)):
        if page_no:
            rows = 0
        lines = []
        for raw_line in page.split(NEWLINE_SEQ):
            if vwf:
                import vwf_metrics
                lines.extend(wrap_line(raw_line, VWF_LINE_PX, measure=vwf_metrics.width))
            else:
                lines.extend(wrap_line(raw_line))
        current = []
        for line in lines:
            if rows >= BOX_ROWS:
                # Box full: this line goes on a fresh page. (If this is the cell's
                # first line, the page break comes right after the previous cell's
                # \n, which only moves the cursor below the box - nothing is drawn.)
                out_pages.append(current)
                current, rows = [], 0
            current.append(line)
            rows += 1
        out_pages.append(current)
    return SPLIT_SEQ.join(NEWLINE_SEQ.join(p) for p in out_pages), rows


_ORIGINAL_SCRIPTS = {}


def cell_ends_page(filename, offset):
    """Does the original script end the text box's page after the text at
    `offset`? A \\n (5C 6E) continues the page; \\f (5C 66) or any script
    opcode ends it."""
    if filename not in _ORIGINAL_SCRIPTS:
        path = os.path.join('original', 'decompressed', filename)
        if not os.path.exists(path):
            path = os.path.join('original', filename)
        _ORIGINAL_SCRIPTS[filename] = open(path, 'rb').read()
    data = _ORIGINAL_SCRIPTS[filename]
    i = offset
    while i < len(data):
        b = data[i]
        if 0x81 <= b <= 0x9f or b >= 0xe0:
            i += 2
            continue
        if b == 0x5c and i + 1 < len(data) and data[i + 1] == 0x6e:
            return False
        if b < 0x20 or b == 0x5c:
            return True
        i += 1
    return True


if __name__ == '__main__':
    OriginalBOD = Disk(SRC_DISK, dump_excel=Dump, pointer_excel=PtrDump)
    TargetBOD = Disk(DEST_DISK)

    # First, extract the current save file from the patched game so I don't lose progress
    # tools/reinsert_without.py resets the patched disk first, so it copies the
    # saves out beforehand and sets this; otherwise we'd copy the blank originals.
    if not os.environ.get('KURO_SAVES_ALREADY_COPIED'):
        TargetBOD.extract('A.FA1', path_in_disk="B-DRKNS", dest_path="patched")
        unpack(b'A.FA1', file_dir=b'patched')

        copyfile(os.path.join('patched', 'BD_FLAG0.DAT'), os.path.join('original', 'BD_FLAG0.DAT'))
        copyfile(os.path.join('patched', 'BD_FLAG1.DAT'), os.path.join('original', 'BD_FLAG1.DAT'))
        copyfile(os.path.join('patched', 'BD_FLAG2.DAT'), os.path.join('original', 'BD_FLAG2.DAT'))
        copyfile(os.path.join('patched', 'BD_FLAG3.DAT'), os.path.join('original', 'BD_FLAG3.DAT'))
        copyfile(os.path.join('patched', 'BD_FLAG4.DAT'), os.path.join('original', 'BD_FLAG4.DAT'))
        copyfile(os.path.join('patched', 'BD_FLAG5.DAT'), os.path.join('original', 'BD_FLAG5.DAT'))
        copyfile(os.path.join('patched', 'BD_FLAG6.DAT'), os.path.join('original', 'BD_FLAG6.DAT'))
        copyfile(os.path.join('patched', 'BD_FLAG7.DAT'), os.path.join('original', 'BD_FLAG7.DAT'))
        copyfile(os.path.join('patched', 'BD_FLAGH.DAT'), os.path.join('original', 'BD_FLAGH.DAT'))

    #input()

    # (otherwise) Fresh start each reinsertion
    #copyfile(SRC_DISK, DEST_DISK)

    # Because the archives get re-inserted with all files, need to copy all the original files into the
    # patched directory to ensure a fresh start.
    for f in FILES:
        if f.source.decode("ASCII") in ARCHIVES_TO_REINSERT:
            filename = f.name.decode("ASCII")
            
            original_path = os.path.join('original', filename)
            patched_path = os.path.join('patched', filename)
            copyfile(original_path, patched_path)

    for filename in FILES_TO_REINSERT:
        #print(filename)
        try:
            original_path = os.path.join('original', 'decompressed', filename)
            patched_path = os.path.join('patched', filename)
            copyfile(original_path, patched_path)

        except FileNotFoundError:
            original_path = os.path.join('original', filename)
            patched_path = os.path.join('patched', filename)
            copyfile(original_path, patched_path)

        # original/decompressed/ holds RAM dumps, and a dump can run past the end of its
        # file into whatever was loaded next: 27KKII01.SCN and 10TNII01.SCN end in ~25
        # bytes of a neighbour's script (10TNII01's tail is someone else's Japanese and
        # print opcodes) where the real file has its trailing table. Building from the
        # dump bakes that in. Start from the real compressed stream instead - the same
        # fix the BSD path below already has. For every other script the two are
        # byte-identical, so this changes nothing there.
        if filename.endswith('.SCN') and os.path.exists(os.path.join('original', filename)):
            dumped = open(patched_path, 'rb').read()
            real = decompress(open(os.path.join('original', filename), 'rb').read())
            if real != dumped and len(real) == len(dumped):
                open(patched_path, 'wb').write(real)
                print('  %s: using the decompressed stream, not the RAM dump (%d bytes differ)'
                      % (filename, sum(a != b for a, b in zip(real, dumped))))

        gf = KuroGamefile(patched_path, disk=OriginalBOD, dest_disk=TargetBOD, pointer_constant=0)

        if filename in UNKNOWN_FORMAT_SCN:
            over = too_long_for_fixed(gf)
            if over:
                # patched/ holds the decompressed original now, which repack recognises
                # as unchanged and stores as the original compressed bytes.
                SKIPPED_TOO_LONG.append((filename, over))
                print('WARNING: %s left in Japanese - %d line(s) longer than the Japanese, '
                      'which this script cannot take yet: %s' % (
                          filename, len(over), ', '.join('%#x %r (max %d)' % (o, e.decode('cp932', 'replace'), n)
                                                         for o, e, n in over[:4])))
                continue

        # TEMP: Let's see if we can get the pointers to be aware of block structure this way
        if filename in FILE_BLOCKS:
            blocks = [Block(gf, (start, stop)) for start, stop in FILE_BLOCKS[filename]]
            gf.blocks = blocks
        else:
            gf.blocks = []

        if filename in POINTERS_TO_REASSIGN:
            #print("Time to reassign some pointers")
            reassignments = POINTERS_TO_REASSIGN[filename]
            for src, dest in reassignments:
                #print("Reassigning", hex(src), hex(dest))
                if src not in gf.pointers:
                    #print("Skipping this one: %s, %s" % (hex(src), hex(dest)))
                    continue
                if dest not in gf.pointers:
                    #print("No pointer for that dest. We'll just move the src pointer to the dest")
                    gf.pointers[dest] = []
                #assert src in gf.pointers
                #assert dest in gf.pointers
                diff = dest - src
                assert dest == src + diff
                for p in gf.pointers[src]:
                    p.edit(diff)
                gf.pointers[dest] += gf.pointers[src]
                gf.pointers.pop(src)

        if filename in BYTE_EDITS:
            #print(BYTE_EDITS[filename])
            for (loc, value) in BYTE_EDITS[filename]:
                gf.edit(loc, value)

        # Every length change, in original offsets: (where, old length, new length). This is
        # the exact original->patched map, which fix_pointers and stale_report use instead of
        # re-deriving it with difflib (which can slip a few bytes and "repair" a pointer that
        # was right - 05SKS03 0x26b).
        edits = []
        too_long = None
        for block in gf.blocks:
            #print(block)
            previous_text_offset = block.start
            diff = 0
            last_text_end = None
            #print(repr(block.blockstring))
            if filename.endswith('SCN'):
                #print(filename)
                translations = translations_of(block, include_blank=True, sheet_name="SCNs")
                #print(translations)
            elif filename.endswith('BSD'):
                #print("Using the BSDs sheet")
                translations = translations_of(block, include_blank=True, sheet_name="BSDs")

            else:
                translations = translations_of(block, include_blank=True)
            page_rows = 0      # rows used on the current text-box page (SCN typesetting)
            prev_t = None
            for t in translations:
                #print(t)
                if filename.endswith('.SCN') and prev_t is not None and cell_ends_page(filename, prev_t.location):
                    page_rows = 0
                prev_t = t
                if t.en_bytestring == b'':
                    # a blank name plate is filled from the glossary, so a name is
                    # written down once rather than in every row that shows it
                    filled = (name_plate_english(t.jp_bytestring)
                              if filename.endswith('.SCN') else None)
                    if filled:
                        t.en_bytestring = filled
                        NAME_FILLED.append((filename, t.location, filled))
                    else:
                        t.en_bytestring = t.jp_bytestring
                # A mapping build writes FILE-INDEX placeholders that already fit the
                # line and carry the original's own control codes, so typesetting them
                # only re-encodes each newline from 2 bytes to 5 and makes the file grow
                # when it is meant to stay exactly the same size.
                typeset_scn = (filename.endswith('.SCN') and filename not in FIXED_LENGTH_SCN
                               and not MAPPING_BUILD)
                if typeset_scn and t.en_bytestring == t.jp_bytestring:
                    # untranslated: its Japanese rows still take up the box
                    page_rows += 1 + t.jp_bytestring.count(b'\\n')
                    typeset_scn = False
                    
                for cc in CONTROL_CODES:
                    if cc in t.en_bytestring:
                        t.en_bytestring = t.en_bytestring.replace(cc, CONTROL_CODES[cc])

                if t.en_bytestring != t.jp_bytestring and filename.endswith('.SCN'):
                    code = script_decode.code_bytes(filename)
                    if any(x in code for x in range(t.location, t.location + len(t.jp_bytestring))):
                        raise ValueError(
                            '%s @%#x is not text: the dumper read script code as Japanese '
                            '(the decoder places an instruction there). Clear its English - '
                            'writing it would overwrite the instruction.' % (filename, t.location))
                    if file_shows_no_text(filename):
                        raise ValueError(
                            '%s @%#x is not text: the game never prints anything from this '
                            'file - not one of its dump rows follows a print instruction. '
                            '31END.SCN is the ending sequence, x86 code rather than script, '
                            'and its four "strings" are stray bytes. Clear its English.'
                            % (filename, t.location))

                if t.en_bytestring != t.jp_bytestring:
                    #print(t.en_bytestring)
                        # Prepend 85, add 1f if it's a num, sub 2 if it's a char
                    if 0xb0 <= t.jp_bytestring[0] <= 0xdf:
                        #print("This is a halfwidth kana string")
                        new_bytestring = b''
                        for b in t.jp_bytestring:
                            if 0xb0 <= b <= 0xdf:
                                new_bytestring += b'\x85' + (b - 2).to_bytes(1, 'little')
                            elif b == 0x20:
                                new_bytestring += b'\x20'
                            else:
                                new_bytestring += b'\x85' + (b + 0x1f).to_bytes(1, 'little')
                        #print(t.jp_bytestring)
                        #print(new_bytestring)
                        t.jp_bytestring = new_bytestring

                    # Auto-indent lines that begin with a quote in SCN files
                    #if t.en_bytestring:
                    #    if t.en_bytestring.startswith(b'"') and filename.endswith('SCN'):
                    #        print(t.en_bytestring[0], t.en_bytestring[0] == 0x20)
                    #        t.en_bytestring = b' ' + t.en_bytestring

                    if filename in FIXED_LENGTH_SCN:
                        if len(t.en_bytestring) > len(t.jp_bytestring):
                            raise ValueError('%s @%s: %r is %d bytes, over the %d-byte Japanese length' % (
                                filename, hex(t.location), t.en_bytestring,
                                len(t.en_bytestring), len(t.jp_bytestring)))
                        t.en_bytestring = t.en_bytestring.ljust(len(t.jp_bytestring), b' ')
                    elif typeset_scn and t.en_bytestring:
                        # Typeset SCN files (word-wrap for text box, new page when full)
                        # With --gfx-dialogue, pixel widths only where the string is known
                        # to print in a box frame (tools/dialogue_modes.py); elsewhere it
                        # may be text mode, where a VWF-length line would overflow.
                        vwf = False
                        if VWF:
                            import dialogue_modes
                            mode = dialogue_modes.mode_of(filename, t.location)
                            vwf = mode == 'g'
                            TYPESET_COUNTS['vwf' if vwf else 'text' if mode else 'unknown'] += 1
                        t.en_bytestring, page_rows = typeset(t.en_bytestring, page_rows, vwf=vwf)
                        typeset_scn = False      # rows already counted


                    #print(t)
                    loc_in_block = t.location - block.start + diff

                    # New style of reinsertion - prevents the issue where you'd need to start
                    # at the first repetition of a string
                    this_original_segment = block.blockstring[loc_in_block:]
                    this_segment = block.blockstring[loc_in_block:]
                    #print([hex(i) for i in t.jp_bytestring])
                    #print([hex(i) for i in this_segment])
                    assert t.jp_bytestring in this_segment, (
                        f'{filename} @{t.location:#x}: the Japanese is no longer where the dump says; '
                        'an earlier pointer edit probably wrote into this text (a false-positive pointer)')
                    idx_in_segment = this_segment.index(t.jp_bytestring)
                    this_segment = this_segment.replace(t.jp_bytestring, t.en_bytestring, 1)
                    block.blockstring = block.blockstring.replace(this_original_segment, this_segment)
                    last_text_end = loc_in_block + idx_in_segment + len(t.en_bytestring)
                    edits.append((t.location + idx_in_segment, len(t.jp_bytestring), len(t.en_bytestring)))


                    # Old style. Getting replaced
                    """
                    #print(t.jp_bytestring)
                    i = block.blockstring.index(t.jp_bytestring)
                    j = block.blockstring.count(t.jp_bytestring)

                    index = 0
                    while index < len(block.blockstring):
                        index = block.blockstring.find(t.jp_bytestring, index)
                        if index == -1:
                            break
                        index += len(t.jp_bytestring) # +2 because len('ll') == 2

                    if j > 1:
                        print("%s multiples of this string found" % j)
                    assert loc_in_block == i, (hex(loc_in_block), hex(i))
                    """

                    #block.blockstring = block.blockstring.replace(t.jp_bytestring, t.en_bytestring, 1)

                if gf.pointers:
                    gf.edit_pointers_in_range((previous_text_offset, t.location), diff)

                previous_text_offset = t.location

                this_diff = len(t.en_bytestring) - len(t.jp_bytestring)
                diff += this_diff

            block_diff = len(block.blockstring) - len(block.original_blockstring)

            # SCN length-sensitive blocks that shrank get padded right after the last
            # translated string (below), which puts everything after it back in its
            # original position - so pointers past that string must not be shifted.
            pads_after_last_text = (filename.endswith('SCN') and
                                    (block.start, block.stop) in LENGTH_SENSITIVE_BLOCKS.get(filename, []) and
                                    block_diff < 0)

            if pads_after_last_text and gf.pointers:
                # Untranslated rows after the last replaced string were already shifted
                # in the loop above, but the padding puts them back.
                print('WARNING: %s is padded after its last string but has pointers - '
                      'pointers targeting untranslated rows after it will be off' % block)

            # 03YSK01A has some pointers pointing to near the end of the file. Need a way to edit those
            if gf.pointers and not pads_after_last_text:
                gf.edit_pointers_in_range((previous_text_offset, block.stop), diff)

            # Ignore size differences in .SCN files
            if filename in LENGTH_SENSITIVE_BLOCKS:
                if (block.start, block.stop) in LENGTH_SENSITIVE_BLOCKS[filename]:
                    print(block, "is length sensitive")
                    if block_diff < 0:
                        print("block_diff of", block, "is", block_diff)
                        padding_len = (-1)*block_diff
                        if filename.endswith('SCN'):
                            # Pad inside the last translated string, before its terminator.
                            # A block's end can fall in script bytecode, where 0x20 is a
                            # real opcode - trailing spaces there crash the interpreter.
                            assert last_text_end is not None, block
                            block.blockstring = (block.blockstring[:last_text_end] +
                                                 b' '*padding_len +
                                                 block.blockstring[last_text_end:])
                            last_at, last_len, _ = max(e for e in edits if block.start <= e[0] < block.stop)
                            edits.append((last_at + last_len, 0, padding_len))
                        else:
                            block.blockstring += padding_len*b'\x00'
                            edits.append((block.stop, 0, padding_len))
                    block_diff = len(block.blockstring) - len(block.original_blockstring)
                    if block_diff > 0:
                        # Grew, and a length-sensitive block can only be padded, not cut.
                        too_long = ('its length-sensitive block %#x-%#x is %d bytes over'
                                    % (block.start, block.stop, block_diff))
                        break
                    assert block_diff == 0, (block_diff, block)

            block.incorporate()
        else:
            too_long = None

        # A script past its RAM slot would load over its neighbour (slot 1 ends where the
        # common script begins), which crashes the game - so it is not inserted at all.
        if not too_long and filename in DECOMPRESSED_SIZE_LIMITS \
                and len(gf.filestring) > DECOMPRESSED_SIZE_LIMITS[filename]:
            too_long = 'it is %d bytes over its %#x-byte RAM slot' % (
                len(gf.filestring) - DECOMPRESSED_SIZE_LIMITS[filename],
                DECOMPRESSED_SIZE_LIMITS[filename])
        if too_long:
            # Leave the original in patched/: repack sees it unchanged and stores the
            # original compressed bytes.
            open(patched_path, 'wb').write(gf.original_filestring)
            SKIPPED_TOO_LONG.append((filename, too_long))
            print('WARNING: %s left in Japanese - the English does not fit: %s. Shorten '
                  'lines in this script (the Story order / Scene columns show which).'
                  % (filename, too_long))
            continue

        # Safety net: a pointer the decoder knows but the pointer sheet lacks still points at
        # the old offset. (Backward pointers used to land here too; pointer_edit.KuroGamefile
        # now places those correctly during the walk.) With the exact edit map, rewrite the
        # ones whose target provably moved while their value did not.
        if filename.endswith('.SCN'):
            fixed, ptr_fixes = fix_pointers.repair(filename, gf.original_filestring, gf.filestring,
                                                   edits=edits)
            if ptr_fixes:
                gf.filestring = fixed
                print('  %s: repaired %d pointer(s) missing from the sheet: %s' % (
                    filename, len(ptr_fixes),
                    ', '.join('@%#x %#x->%#x' % (loc, old, new) for loc, old, new in ptr_fixes)))

        gf.write(skip_disk=True)
        if filename.endswith('.SCN'):
            fix_pointers.save_edits(os.path.join('patched', filename), edits)

        # Check decompressed file size against its RAM slot (see rominfo.SCN_SLOT_SIZES)
        if filename in DECOMPRESSED_SIZE_LIMITS:
            file_size = len(gf.filestring)
            limit = DECOMPRESSED_SIZE_LIMITS[filename]
            if file_size > limit:
                print('WARNING: %s is %s bytes, exceeds RAM limit of %s (+%d bytes over)' % (
                    filename, hex(file_size), hex(limit), file_size - limit))
            else:
                slack = limit - file_size
                print('  %s: %s / %s (%d bytes of slack)' % (
                    filename, hex(file_size), hex(limit), slack))

    # BSD battle script reinsertion using structural disassembly/reassembly
    for filename in BSD_FILES_WITH_TEXT:
        # Build from the real compressed stream, not original/decompressed/: those are
        # RAM dumps, and the game relocates BSD data in place (header, event table,
        # handler variables), so several dumps hold runtime values that must not be
        # baked into the reinserted file.
        try:
            original_compressed = open(os.path.join('original', filename), 'rb').read()
        except FileNotFoundError:
            print('WARNING: %s not found in original/, skipping' % filename)
            continue
        data = decompress(original_compressed)
        patched_path = os.path.join('patched', filename)
        open(patched_path, 'wb').write(data)
        parsed = parse_bsd(data)

        if not parsed['text_regions']:
            continue

        # Build a single block spanning all text for DumpExcel lookup
        gf = Gamefile(patched_path, disk=OriginalBOD, dest_disk=TargetBOD, pointer_constant=0)
        first_text = min(r.offset for r in parsed['text_regions'])
        last_text = max(r.end for r in parsed['text_regions'])
        block = Block(gf, (first_text, last_text + 1))

        translations_raw = translations_of(block, include_blank=True, sheet_name="BSDs")

        # Build translations list for bsd_tool: (offset, jp_bytes, en_bytes)
        translations = []
        for t in translations_raw:
            if t.en_bytestring == b'':
                t.en_bytestring = t.jp_bytestring

            for cc in CONTROL_CODES:
                if cc in t.en_bytestring:
                    t.en_bytestring = t.en_bytestring.replace(cc, CONTROL_CODES[cc])

            # Typeset BSD dialogue text (same word-wrapping as SCN)
            # (no page tracking here yet: BSD pages aren't split automatically)
            if t.en_bytestring and t.en_bytestring != t.jp_bytestring:
                t.en_bytestring, _ = typeset(t.en_bytestring)

            translations.append((t.location, t.jp_bytestring, t.en_bytestring))

        result = reinsert_bsd(data, translations)

        if result != data:
            open(patched_path, 'wb').write(result)
            print('  %s: %d -> %d bytes (%+d)' % (
                filename, len(data), len(result), len(result) - len(data)))
            if len(result) > BSD_DECOMPRESSED_SIZE_LIMIT:
                print('  WARNING: %s is %s bytes, exceeds shared battle-buffer '
                      'limit of %s (+%d bytes over) - see BSD_DECOMPRESSED_SIZE_LIMIT '
                      'in rominfo.py' % (
                          filename, hex(len(result)), hex(BSD_DECOMPRESSED_SIZE_LIMIT),
                          len(result) - BSD_DECOMPRESSED_SIZE_LIMIT))
        else:
            print('  %s: no changes' % filename)

    # Editing compressed files without decompressing them. A truly barbaric practice
    for filename in COMPRESSED_FILES_TO_EDIT:
        print(filename)
        original_path = os.path.join('original', filename)
        patched_path = os.path.join('patched', filename)
        copyfile(original_path, patched_path)

        gf = Gamefile(patched_path, disk=OriginalBOD, dest_disk=TargetBOD, pointer_constant=0)

        if filename in BYTE_EDITS:
            print(BYTE_EDITS[filename])
            for (loc, value) in BYTE_EDITS[filename]:
                gf.edit(loc, value)
                
        gf.write(skip_disk=True)

    # Dialogue boxes in MB3N's graphics mode (python reinsert.py --gfx-dialogue): the
    # groundwork for a variable-width font, see docs/vwf_findings.md. Off by default: it
    # looks the same as text mode, and the text-layer readback that verify_text.py,
    # verify_blocks.py and the bench use sees no dialogue in graphics mode.
    code_patched = []
    if '--gfx-dialogue' in sys.argv:
        sys.path.insert(0, 'tools')
        import mb3_gfx          # needs keystone (pip install keystone-engine)
        import dialogue_gfx
        for name, patch in (('MB3N.BIN', mb3_gfx.patch), ('99CMN.SCN', dialogue_gfx.patch)):
            if name in FILES_TO_REINSERT:
                raise SystemExit('%s is now reinserted from the workbook; the graphics-mode '
                                 'patch expects the original file and needs updating' % name)
            data = decompress(open(os.path.join('original', name), 'rb').read())
            with open(os.path.join('patched', name), 'wb') as fh:
                fh.write(patch(data))
            code_patched.append(name)
            print('  %s: graphics-mode dialogue patch applied' % name)
        print('  typeset: %(vwf)d cells by pixel width (VWF), %(text)d by characters (text '
              'mode), %(unknown)d by characters (never reached by the mode sweep)' % TYPESET_COUNTS)
        # BD.BIN is already reinserted (its text), so patch the reinserted copy
        bd_path = os.path.join('patched', 'BD.BIN')
        with open(bd_path, 'rb') as fh:
            bd = fh.read()
        with open(bd_path, 'wb') as fh:
            fh.write(dialogue_gfx.patch_bd(bd))
        print('  BD.BIN: back to text mode at the end of every script run')

    for filename in ARCHIVES_TO_REINSERT:
        gamefile_path = os.path.join('patched', filename)

        repack(gamefile_path, also_reinserted=code_patched)
        # Gotta repack it first, then initialize the gamefile
        gf = Gamefile(gamefile_path, disk=OriginalBOD, dest_disk=TargetBOD)

        gf.write(path_in_disk='B-DRKNS')