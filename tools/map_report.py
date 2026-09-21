"""Merge every autoplay run's map into one picture of the game.

Each run writes routes/<run>_map.json: what each NPC and trigger zone on each map did,
and which dump rows its text came from. This merges them all and adds the bit a single
run can't know - how much of the workbook has actually been seen on screen.

    python tools/map_report.py [--out docs/game_map.md]
"""
import argparse
import collections
import glob
import json
import os

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # the project root; this file is in tools/


def load():
    graph, coverage, runs = {}, collections.defaultdict(dict), []
    map_files = collections.defaultdict(set)
    for path in sorted(glob.glob(os.path.join(HERE, 'routes', '*_map.json'))):
        run = os.path.basename(path)[:-len('_map.json')]
        try:
            d = json.load(open(path))
        except ValueError:
            continue
        runs.append(run)
        for key, g in d.get('graph', {}).items():
            old = graph.get(key)
            if old is None or len(g.get('text', [])) > len(old.get('text', [])):
                graph[key] = g
        for fn, offs in d.get('coverage', {}).items():
            for off, maps in offs.items():
                coverage[fn].setdefault(off, set()).update(maps)
        for m, files in d.get('map_files', {}).items():
            map_files[m].update(files)
    return graph, coverage, runs, map_files


def workbook_rows():
    """filename -> {offset: has_english}, so coverage can be given as a fraction."""
    import openpyxl
    from rominfo import DUMP_XLS_PATH
    out = collections.defaultdict(dict)
    wb = openpyxl.load_workbook(DUMP_XLS_PATH, read_only=True)
    for sheet in wb.sheetnames:
        for row in wb[sheet].iter_rows(min_row=2, values_only=True):
            fn, off, jp, jl, en = (list(row) + [None] * 5)[:5]
            if fn and off:
                out[fn][str(off)] = isinstance(en, str) and bool(en.strip())
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default=os.path.join('docs', 'game_map.md'))
    args = ap.parse_args()
    graph, coverage, runs, map_files = load()
    rows = workbook_rows()
    by_map = collections.defaultdict(list)
    for g in graph.values():
        by_map[g['map']].append(g)
    os.makedirs(os.path.dirname(args.out) or '.', exist_ok=True)
    with open(args.out, 'w', encoding='utf-8') as f:
        def out(*t, **kw):
            print(*t, file=f, **kw)

        out('# Kuro no Ken: what the autoplayer has mapped', end='\n\n')
        out('Merged from %d run(s): %s.' % (len(runs), ', '.join(runs)), end='\n\n')
        out('Each entry is an NPC or trigger zone that was tried, where it led, and the')
        out('dump rows its text came from (`file offset`, English or Japanese column).')
        out('docs/script_map.md covers the same ground statically, for the whole game.',
            end='\n\n')
        for m in sorted(by_map):
            out('## %s' % m, end='\n\n')
            if map_files.get(m):
                out('Scripts loaded here: %s (docs/script_map.md lists their dump rows)'
                    % ', '.join(sorted(map_files[m])), end='\n\n')
            for g in sorted(by_map[m], key=lambda g: (g['kind'], g['index'])):
                bits = ['%-4s %2d' % (g['kind'], g['index'])]
                if g.get('leads_to'):
                    bits.append('-> %s' % g['leads_to'])
                if g.get('pages'):
                    bits.append('%d page(s)' % g['pages'])
                if g.get('outcome') not in (None, 'ok'):
                    bits.append('[%s]' % str(g['outcome'])[:50])
                srcs = sorted({'%s %s' % (a, b) for a, b in g.get('text', [])})
                if srcs:
                    bits.append('text: ' + ', '.join(srcs[:8]) + ('...' if len(srcs) > 8 else ''))
                out('- ' + '  '.join(bits))
            out()
        out('## Dump coverage', end='\n\n')
        out('How much of each file has been seen on screen. "translated" counts rows with')
        out('English in the workbook, so a file at 0/N translated but seen in play is text')
        out('that still needs translating.', end='\n\n')
        out('| file | rows seen | rows in workbook | translated |')
        out('|---|---|---|---|')
        seen_total = 0
        for fn in sorted(coverage):
            n = len(coverage[fn])
            seen_total += n
            have = rows.get(fn, {})
            out('| %s | %d | %d | %d |' % (fn, n, len(have), sum(1 for v in have.values() if v)))
        out()
        out('%d rows seen across %d files; the workbook has %d rows in %d files.'
            % (seen_total, len(coverage), sum(len(v) for v in rows.values()), len(rows)))
        unseen = [fn for fn in sorted(rows) if fn not in coverage]
        out(end='\n')
        out('### Files never reached yet (%d)' % len(unseen), end='\n\n')
        out(', '.join('`%s`' % fn for fn in unseen))
    print('%d maps, %d entries, %d files with seen text -> %s'
          % (len(by_map), len(graph), len(coverage), args.out))


if __name__ == '__main__':
    main()
