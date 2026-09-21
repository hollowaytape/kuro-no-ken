"""Cells whose English runs on into the next cell.

The game shows consecutive cells of a file in one text box, but `reinsert.typeset` wraps
each cell on its own. When a translated sentence is split across two cells, the first
one's tail gets a line to itself and the second starts on the next line - on screen that
is an orphan word like

    I was prepared to receive the full wrath of
    the
    Grand Duke's over the loss of the Starlight

and, if the box was nearly full, the remainder is pushed onto a page of its own
("Emerald..." alone). Neither shows up when a cell is rendered by itself, which is why
verify_text.py passes these.

    python tools/split_sentences.py [--limit N]
"""
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..'))
import openpyxl                                  # noqa: E402
from rominfo import DUMP_XLS_PATH                # noqa: E402

# A cell that ends mid-sentence: no terminator, no closing quote, not a control code.
ENDS_OPEN = re.compile(r'[A-Za-z,;:\-]\s*$')
CODE = re.compile(r'\\[a-z](?:\d+(?:,\d+)*)?|\[SPLIT\]')

wb = openpyxl.load_workbook(DUMP_XLS_PATH, read_only=True)
rows = []
for sheet in wb.sheetnames:
    for r in wb[sheet].iter_rows(min_row=2, values_only=True):
        fn, off, jp, jl, en = (list(r) + [None] * 5)[:5]
        if fn and isinstance(off, str) and isinstance(en, str) and en.strip():
            try:
                rows.append((fn, int(off, 16), en))
            except ValueError:
                pass
rows.sort(key=lambda x: (x[0], x[1]))

hits = []
for (fn, off, en), (fn2, off2, en2) in zip(rows, rows[1:]):
    if fn != fn2:
        continue
    tail = CODE.sub('', en).rstrip()
    if not tail or not ENDS_OPEN.search(tail):
        continue
    head = CODE.sub('', en2).lstrip()
    if not head or not head[0].isalpha() or head[0].isupper() and not tail.endswith(','):
        # a new sentence starting with a capital is fine unless the previous cell
        # ended on a comma or a dangling word
        if not tail.endswith((',', '-', ';', ':')) and len(tail.split()[-1]) > 3:
            continue
    hits.append((fn, off, tail.split()[-1], off2, head[:40]))

# How bad it looks depends on the last wrapped line of the first cell: the game breaks
# where the Japanese did, so that tail gets a line of its own. One or two words is an
# obvious orphan; a full line reads fine.
from reinsert import wrap_line                   # noqa: E402

ranked = []
for fn, off, word, off2, head in hits:
    cell = next(en for f, o, en in rows if f == fn and o == off)
    lines = wrap_line(CODE.sub('', cell).encode('cp932', 'replace'))
    tail = lines[-1].decode('cp932', 'replace').strip() if lines else ''
    ranked.append((len(tail), fn, off, tail, off2, head))
ranked.sort()

limit = int(sys.argv[sys.argv.index('--limit') + 1]) if '--limit' in sys.argv else 30
print('%d cell pairs where the English runs on into the next cell' % len(hits))
print('worst first - the last line of the first cell is what the player sees alone:')
for n, fn, off, tail, off2, head in ranked[:limit]:
    print('   %-14s %#07x last line %-22r -> %#07x "%s"' % (fn, off, tail, off2, head[:34]))
