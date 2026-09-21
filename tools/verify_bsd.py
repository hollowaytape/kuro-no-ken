"""Static verification for the battle scripts (.BSD).

The 47 BSD files that hold text are the one part of the script with no verification
at all: `verify_blocks.py` runs .SCN blocks in the emulator, but a battle script is
driven by the battle engine, not the interpreter, so there is nothing to call. This
checks the parts that can be checked without running the game, and they are the parts
that actually break:

1. **round trip** - reassembling a file with no changes must reproduce it byte for byte.
   If it does not, `bsd_tool` misparses the file and reinsertion would corrupt it.
2. **relocation** - grow every line and check every fixup still points at the same
   bytes it pointed at before. A BSD is full of absolute offsets (resource table, 16
   event entry points, `cs:` stores, near branches), so this is where a translated
   battle goes wrong.
3. **reach** - every dump row's Japanese must land inside a parsed text region.
   A row outside one is a row reinsert silently drops.

    python tools/verify_bsd.py                 # all files with dumped text
    python tools/verify_bsd.py --all           # all 219, including the ones with no text
    python tools/verify_bsd.py --file D010_X10.BSD -v
"""
import argparse
import collections
import glob
import os
import sys

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # the project root; this file is in tools/
sys.path.insert(0, HERE)

import bsd_tool                                     # noqa: E402

ORIG = os.path.join(HERE, 'original', 'decompressed')
WINDOW = 12            # bytes compared at each relocated target


def dump_rows():
    """filename -> [(offset, japanese, english)] for every BSD row in the workbook."""
    import openpyxl
    from rominfo import DUMP_XLS_PATH
    rows = collections.defaultdict(list)
    wb = openpyxl.load_workbook(DUMP_XLS_PATH, read_only=True)
    for sheet in wb.sheetnames:
        for r in wb[sheet].iter_rows(min_row=2, values_only=True):
            fn, off, jp = r[0], r[1], r[2]
            en = r[4] if len(r) > 4 else None
            if not (fn and str(fn).endswith('.BSD') and isinstance(jp, str) and jp.strip()):
                continue
            try:
                at = int(str(off), 16)
            except ValueError:
                continue
            rows[str(fn)].append((at, jp, en))
    return rows


def check_file(fn, rows, verbose=False):
    """Returns (list of problems, stats dict)."""
    data = open(os.path.join(ORIG, fn), 'rb').read()
    bad = []
    parsed = bsd_tool.parse_bsd(data)
    regions = parsed['text_regions']
    fixups = parsed['fixups']

    # 1. round trip
    same = bsd_tool.reassemble(parsed)
    if same != data:
        where = next((i for i in range(min(len(same), len(data))) if same[i] != data[i]), -1)
        bad.append('round trip differs at %#x (len %d -> %d)' % (where, len(data), len(same)))
        return bad, dict(regions=len(regions), fixups=len(fixups), rows=len(rows), unreached=0)

    # 2. relocation under growth
    if regions:
        grown = {i: r.text_bytes + b'.' * 8 for i, r in enumerate(regions)}
        out = bsd_tool.reassemble(parsed, grown)
        changes = sorted((r.end, 8) for r in regions)

        def moved(off):
            return off + bsd_tool.compute_shift(off, changes)

        text_span = [(r.offset, r.end) for r in regions]

        def in_text(o):
            return any(a <= o < b for a, b in text_span)

        # Bytes that are themselves offsets: they are *supposed* to read differently
        # after relocation, so they are skipped when comparing what a pointer points at.
        reloc = set()
        for f in fixups:
            reloc.update(range(f.file_offset, f.file_offset + f.size))

        import struct
        for f in fixups:
            loc = moved(f.file_offset)
            if f.size != 2 or loc + 2 > len(out):
                continue
            got = struct.unpack_from('<H', out, loc)[0]
            if f.kind == 'cs_mov_target':
                # not a pointer into the file: the game uses the first byte past the
                # loaded data as scratch, so this always holds the new file length.
                want = len(out)
            elif f.kind == 'rel16':
                # a near branch stores a displacement from the end of the instruction,
                # so both ends move and the stored word only changes when they move
                # by different amounts.
                want = (moved(f.value) - (loc + 2)) & 0xffff
            else:
                want = moved(f.value)
            if got != want:
                bad.append('%s at %#x: holds %#x, expected %#x (was %#x, target %#x)'
                           % (f.kind, f.file_offset, got, want,
                              struct.unpack_from('<H', data, f.file_offset)[0], f.value))
                continue
            if f.kind in ('rel16', 'cs_mov_target') or in_text(f.value):
                continue     # target is the text itself; its bytes are meant to differ
            if f.value >= len(data):
                continue     # points just past the file - scratch space, nothing to compare
            tgt = moved(f.value)
            for k in range(WINDOW):
                if f.value + k in reloc or f.value + k >= len(data) or tgt + k >= len(out):
                    continue
                if data[f.value + k] != out[tgt + k]:
                    bad.append('%s at %#x: target %#x -> %#x differs at +%d (%02x vs %02x)'
                               % (f.kind, f.file_offset, f.value, tgt, k,
                                  data[f.value + k], out[tgt + k]))
                    break
        want_len = len(data) + 8 * len(regions)
        if len(out) != want_len:
            bad.append('grown file is %d bytes, expected %d' % (len(out), want_len))

    # 3. every dump row inside a text region
    unreached = 0
    for at, jp, _en in rows:
        want = jp.encode('cp932', 'replace')
        hit = next((r for r in regions if r.offset - 4 <= at < r.end), None)
        if hit is None or want not in hit.text_bytes:
            unreached += 1
            if verbose:
                bad.append('row %#x %r is not inside any text region' % (at, jp[:24]))
    if unreached and not verbose:
        bad.append('%d of %d dump row(s) fall outside a parsed text region'
                   % (unreached, len(rows)))
    return bad, dict(regions=len(regions), fixups=len(fixups), rows=len(rows),
                     unreached=unreached)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--file')
    ap.add_argument('--all', action='store_true', help='include BSDs with no dumped text')
    ap.add_argument('-v', '--verbose', action='store_true')
    args = ap.parse_args()

    rows = dump_rows()
    names = ([args.file] if args.file else
             sorted(os.path.basename(p) for p in glob.glob(os.path.join(ORIG, '*.BSD'))))
    if not args.all and not args.file:
        names = [n for n in names if n in rows]

    tot = collections.Counter()
    failed = []
    for fn in names:
        try:
            bad, st = check_file(fn, rows.get(fn, []), args.verbose)
        except Exception as ex:
            bad, st = ['%s: %s' % (type(ex).__name__, ex)], dict(regions=0, fixups=0,
                                                                rows=0, unreached=0)
        for k, v in st.items():
            tot[k] += v
        tot['files'] += 1
        if bad:
            failed.append(fn)
            print('FAIL  %-14s %d region(s), %d fixup(s)' % (fn, st['regions'], st['fixups']))
            for b in bad[:8 if not args.verbose else 100]:
                print('        %s' % b)
        elif args.verbose or args.file:
            print('ok    %-14s %d region(s), %d fixup(s), %d row(s)'
                  % (fn, st['regions'], st['fixups'], st['rows']))

    print()
    print('%d file(s): %d text region(s), %d fixup(s), %d dump row(s)'
          % (tot['files'], tot['regions'], tot['fixups'], tot['rows']))
    print('%d file(s) failed, %d row(s) outside a text region'
          % (len(failed), tot['unreached']))
    return 1 if failed else 0


if __name__ == '__main__':
    sys.exit(main())
