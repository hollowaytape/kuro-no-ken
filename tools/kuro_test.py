"""Automated play-testing for Kuro no Ken: read and check every dialogue page,
replay recorded routes, and record new routes while a human plays.

    python tools/kuro_test.py lint                 # static check of the workbook: overfull pages, smart quotes
    python tools/kuro_test.py explore              # from here: every NPC and zone on this map, rolling back between each
    python tools/kuro_test.py read                 # page through the open dialogue, checking it
    python tools/kuro_test.py record routes/x.json # log your own play into a route file (Ctrl+C to stop)
    python tools/kuro_test.py play routes/x.json   # replay a route, checking all text; writes a report

Checks on every dialogue page (see check_page):
  - crash      : the DOS prompt (or an error message) replaced the game
  - untranslated : Japanese on screen; reports the workbook file/offset if found
  - not-inserted : Japanese on screen although the workbook has English for it
  - nonascii   : other full-width text, e.g. a typographic quote/dash in the English
                 (renders as a wide glyph)
  - garbage    : cells that aren't printable characters or the game's glyphs
  - codes      : a leftover control code ("\\n", "[BR]"...) printed literally
  - overflow   : a line reaching the prompt icon at the box's right edge
  - truncated  : the page ends mid-sentence: the box ran out of rows (name + 4 lines)
                 and the rest of the cell was never shown
  - unmatched  : the page text isn't in the English script at all, which usually
                 means a bad pointer is showing the wrong (or a cut-off) string
Each matched page also reports which file/offset it came from.
"""
import json
import os
import re
import sys
import time
from difflib import SequenceMatcher

from kuro_emu import Emu, Blocked, KEYS, ICON_MARGIN, GLYPH, BAD, MAP_SEG_PTR

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # the project root; this file is in tools/
REPORT_DIR = os.path.join(HERE, 'test_reports')

_CODE_RE = re.compile(r'\[[A-Za-z0-9]+\]|\\[a-z](?:\d+(?:,\d+)?)?')
_WS_RE = re.compile(r'\s+')


_QUOTES = str.maketrans({'‘': "'", '’': "'", '“': '"', '”': '"'})
_SPLIT_RE = re.compile(r'\[SPLIT\]|\\f')


def same_map(a, b):
    """16d8:0951 is the last-loaded resource (fld1, fld1.mp1, fld1.mpc...), so
    compare map names without the extension."""
    return a.split('.')[0].lower() == b.split('.')[0].lower()


_JP_STRIP = str.maketrans('', '', '「」　 ・')   # 「」, full-width space, space, ・


def _jp_norm(s):
    """Japanese text for lookup: drop control codes, brackets and spacing."""
    return _CODE_RE.sub('', s).translate(_JP_STRIP)


def normalize(s):
    s = _CODE_RE.sub(' ', s.translate(_QUOTES))
    return _WS_RE.sub(' ', s).strip().lower()


class Script:
    """The English script from the dump workbook, as one normalized string per
    source file (so a page that spans several cells still matches as a substring)."""

    def __init__(self, path=os.path.join(HERE, os.environ.get('KURO_DUMP_XLS', 'KuroNoKen_dump.xlsx'))):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        self.files = {}      # filename -> (normalized text, [(char_index, offset)])
        self.lines = []      # (normalized line, filename, offset) for fuzzy lookup
        self.ends = {}       # filename -> char indexes where a page may end
        self.japanese = []   # (normalized JP, filename, offset, has_english)
        for sheet in wb.sheetnames:
            for row in wb[sheet].iter_rows(min_row=2, values_only=True):
                fn, off, jp, en = row[0], row[1], row[2], row[4]
                if isinstance(jp, str) and jp.strip():
                    has_en = isinstance(en, str) and bool(en.strip()) and not en.startswith('=')
                    self.japanese.append((_jp_norm(jp), fn, off, has_en))
                if not isinstance(en, str) or not en.strip() or en.startswith('='):
                    continue
                n = normalize(en)
                if not n:
                    continue
                text, index = self.files.setdefault(fn, ('', []))
                start = len(text) + (1 if text else 0)
                index.append((start, off))
                text = (text + ' ' + n) if text else n
                self.files[fn] = (text, index)
                # a page may legitimately end at the end of a cell or at a [SPLIT]
                ends = self.ends.setdefault(fn, set())
                pos = start
                for part in _SPLIT_RE.split(en)[:-1]:
                    pos += len(normalize(part)) + 1
                    ends.add(pos - 1)
                ends.add(len(text))
                self.lines.append((n, fn, off))

    def find(self, page_text):
        """Return (filename, offset) of the first script location containing page_text."""
        n = normalize(page_text)
        if not n:
            return None
        for fn, (text, index) in self.files.items():
            i = text.find(n)
            if i >= 0:
                off = max((o for start, o in index if start <= i), default=index[0][1],
                          key=lambda o: int(o, 16) if isinstance(o, str) else o)
                return fn, off
        return None

    def missing_tail(self, page_text):
        """If page_text ends partway through a script cell (not at a [SPLIT]), return
        the rest of that cell - text the box never showed. Otherwise None."""
        n = normalize(page_text)
        for fn, (text, index) in self.files.items():
            i = text.find(n)
            if i < 0:
                continue
            end = i + len(n)
            if end in self.ends[fn] or text[end:end + 1] != ' ':
                return None
            nxt = min((e for e in self.ends[fn] if e > end), default=len(text))
            return text[end:nxt].strip()
        return None

    def find_japanese(self, line, prefer=None):
        """Where a line of on-screen Japanese comes from: (filename, offset,
        has_english) or None. `prefer` = the page's (file, offset): the nearest
        occurrence in that file wins."""
        n = _jp_norm(line)
        if not n:
            return None
        hits = [(fn, off, has_en) for jp, fn, off, has_en in self.japanese if n in jp]
        if not hits:
            # The engine builds some lines at runtime: the item-pickup message is the
            # item's name followed by the script's "を手にいれた。" (99CMN 0xc6c), so
            # the workbook row is a piece of the screen line rather than the reverse.
            hits = [(fn, off, has_en) for jp, fn, off, has_en in self.japanese
                    if len(jp) >= 4 and jp in n]
        if not hits:
            return None
        if prefer:
            same = [h for h in hits if h[0] == prefer[0]]
            if same:   # the occurrence nearest the page's own text
                return min(same, key=lambda h: abs(int(h[1], 16) - int(prefer[1], 16)))
        # A short or common line matches all over the workbook - カイエス (a name) is in
        # 542 rows - and picking the first one invents a source, which then gets
        # reported as "has English but shows Japanese" against an innocent file.
        if len({h[0] for h in hits}) > 3:
            return None
        return hits[0]

    def closest(self, line):
        n = normalize(line)
        best = (0, None)
        for text, fn, off in self.lines:
            if abs(len(text) - len(n)) > max(10, len(n)):
                continue
            r = SequenceMatcher(None, n, text).ratio()
            if r > best[0]:
                best = (r, (text, fn, off))
        return best


# The title screen. A block run out of context can end the game, and then the menu gets
# checked as if it were dialogue - six "unmatched" findings that were nothing of the kind.
TITLE_ITEMS = ('watch opening', 'new game', 'continue', 'color settings', 'exit to dos')


def looks_like_panel(lines):
    """Is this the game's own UI rather than script text?

    Shops and status windows share the text layer with dialogue, so they arrive here
    looking like a page - and then every price, HP readout and "PAGE 1/1" is reported
    as unmatched or non-ASCII, which buries the real findings.
    """
    if any(re.search(r'PAGE\s*\d+\s*/\s*\d+', l) for l in lines):
        return True
    if sum(1 for l in lines if l.strip().lower() in TITLE_ITEMS) >= 2:
        return True
    # A readout line is short and has a number in it: "HP  229", "/  229", "430 G".
    # Counting only lines with two numbers missed the status window, whose every line
    # has one, and its contents then get reported as text that isn't in the script.
    readout = sum(1 for l in lines
                  if re.search(r'\d', l) and (len(l.split()) <= 3 or len(re.findall(r'\d+', l)) >= 2))
    return len(lines) >= 3 and readout >= len(lines) / 2


def check_page(page, script, speakers=(), prefer=None):
    """Return (source, [problems]) for one dialogue page from Emu.dialogue()."""
    problems = []
    lines = [l for l in page['lines'] if l.strip()]
    if looks_like_panel(lines):
        return None, []
    text = ' '.join(l.strip() for l in lines)
    if BAD in text:
        problems.append(('garbage', 'undecodable cells: ' + text))
    jp_lines = [l for l in lines if any(ord(ch) > 0x7f and ch not in (GLYPH, BAD) for ch in l)]
    if re.search(r'\\[a-z]|\[[A-Z][A-Za-z0-9]*\]', text):
        problems.append(('codes', text))
    icon = page.get('icon_col')
    if icon is not None:
        for line, (first, last) in zip(page['lines'], page['cols']):
            if last >= icon - ICON_MARGIN and line.strip():
                problems.append(('overflow', f'col {last} (prompt icon at {icon}): {line.strip()}'))
    source = None
    ascii_lines = [l.translate(_QUOTES) for l in lines
                   if all(ord(ch) < 0x80 for ch in l.translate(_QUOTES))]
    if ascii_lines and not any(p[0] == 'garbage' for p in problems):
        # A speaker's name may share the box; match with and without the first line.
        body = ' '.join(l.strip() for l in ascii_lines)
        source = script.find(body)
        if not source and len(ascii_lines) > 1:
            body = ' '.join(l.strip() for l in ascii_lines[1:])
            source = script.find(body)
        if source:
            tail = script.missing_tail(body)
            if tail:
                problems.append(('truncated', f'box ended before "{tail}" ({len(page["lines"])} rows shown)'))
        if not source:
            bad = [l for l in ascii_lines if not script.find(l.strip())]
            for l in bad:
                ratio, hit = script.closest(l)
                near = f' (closest {ratio:.0%}: {hit[1]} {hit[2]} "{hit[0]}")' if hit else ''
                problems.append(('unmatched', l.strip() + near))
            if not bad:
                # every line exists somewhere but not in this order
                problems.append(('unmatched', 'lines found separately, not as one passage: ' + body))
    for line in jp_lines:
        # `prefer` is the file this text is known to come from (the block runner knows
        # it); without it a line that appears in several scripts cannot be attributed.
        hit = script.find_japanese(line.strip(), prefer=source or prefer)
        if hit and not hit[2]:
            problems.append(('untranslated', f'{hit[0]} {hit[1]}: {line.strip()}'))
        elif hit:
            # the workbook has English for this, yet the game shows the Japanese
            problems.append(('not-inserted', f'{hit[0]} {hit[1]} has English but shows: {line.strip()}'))
        elif re.search('[぀-ヿ一-鿿]', line):
            problems.append(('untranslated', f'(not in workbook) {line.strip()}'))
        else:
            problems.append(('nonascii', line.strip()))
    return source, problems


class GameOver(Exception):
    """The party was wiped out; the route can't continue from here."""


class Tester:
    def __init__(self, emu=None, script=None, log=print, god=True):
        self.god = god         # keep HP topped up in battles (cheat; testing text, not balance)
        self.emu = emu or Emu()
        self.script = script or Script()
        self.log = log
        self.issues = []       # (kind, detail, context)
        self.pages = []        # every page read: (map, source, lines)

    def issue(self, kind, detail):
        ctx = f'{self.emu.map_name()} {tuple(hex(v) for v in self.emu.pos())}'
        self.issues.append((kind, detail, ctx))
        self.log(f'  !! {kind}: {detail}  [{ctx}]')
        if kind in ('crash', 'garbage', 'unmatched', 'overflow', 'codes', 'truncated'):
            os.makedirs(REPORT_DIR, exist_ok=True)
            self.emu.shot(os.path.join(REPORT_DIR, f'issue_{len(self.issues):03d}_{kind}.png'))

    def check_crash(self):
        if self.emu.state() == 'dos':
            self.issue('crash', ' / '.join(r.strip() for r in self.emu.screen_text() if r.strip()))
            return True
        return False

    def wait_page(self, timeout=6.0):
        """Wait for a finished page (prompt icon showing). None if the box closed."""
        end = self.emu.clock() + timeout
        last = None
        while self.emu.clock() < end:
            d = self.emu.dialogue()
            if d and d['waiting']:
                return d
            if d is None and last is None and self.emu.clock() > end - timeout + 1.0:
                return None
            last = d
            self.emu.wait(0.1)
        return self.emu.dialogue()

    def read_dialogue(self, max_pages=200, match=True):
        """Page through the open dialogue with Space, checking each page. Returns
        the number of pages read; stops when the box closes or the game crashes."""
        n = 0
        seen_blank = 0
        last, repeats = None, 0
        pending_trunc = None      # (detail, missing text) awaiting the next page
        # A shop/inventory screen is a menu, not dialogue: pressing Space there buys
        # things and counts a quantity up (autoplay once "read" 199 pages of
        # "Half Plate 120 x 0,1,2..."). Back out instead.
        if self.emu.menu() and self.emu.state() != 'battle':
            self.log('  (menu, not dialogue: backing out)')
            for _ in range(3):
                self.emu.press('ESC', gap=0.3)
                if not self.emu.menu():
                    break
            return 0
        while n < max_pages:
            if self.check_crash():
                return n
            d = self.wait_page()
            if d is not None and any(l.strip() for l in d['lines']):
                if d['lines'] == last:
                    repeats += 1
                    if repeats >= 3:
                        # Space isn't advancing it: not really a dialogue, or input is stuck
                        self.issue('stuck-text', ' / '.join(l.strip() for l in d['lines']))
                        return n
                    self.emu.press('SPACE', gap=0.3)
                    continue
                last, repeats = d['lines'], 0
            if d is None:
                seen_blank += 1
                if seen_blank >= 2:
                    if pending_trunc:
                        self.issue('truncated', pending_trunc[0])
                    return n
                self.emu.wait(0.4)
                continue
            seen_blank = 0
            if d['lines'] and any('Annihilated' in l for l in d['lines']):
                self.issue('gameover', ' / '.join(l.strip() for l in d['lines']))
                raise GameOver()
            if d['lines']:
                source, problems = check_page(d, self.script)
                if not match:    # e.g. battle results, assembled from pieces at runtime
                    problems = [p for p in problems if p[0] != 'unmatched']
                # A 'truncated' page is fine if the missing words open the next page:
                # reinsert.typeset() splits overfull boxes at line boundaries, which the
                # workbook's own page breaks don't show.
                page_text = normalize(' '.join(l.strip() for l in d['lines']))
                if pending_trunc and pending_trunc[1] not in page_text[:len(pending_trunc[1]) + 40]:
                    self.issue('truncated', pending_trunc[0])
                pending_trunc = None
                for kind, detail in list(problems):
                    if kind == 'truncated':
                        tail = re.search(r'before "(.*)" \(', detail)
                        pending_trunc = (detail, normalize(tail.group(1))[:30] if tail else '')
                        problems.remove((kind, detail))
                self.pages.append((self.emu.map_name(), source, d['lines']))
                where = f'{source[0]} {source[1]}' if source else '?'
                self.log(f'  [{where}] ' + ' / '.join(l.strip() for l in d['lines']))
                for kind, detail in problems:
                    self.issue(kind, detail)
                n += 1
            before = self.emu.text_cells()
            self.emu.press('SPACE', gap=0.05)
            # wait for the page to change (the typewriter effect starts immediately)
            end = self.emu.clock() + 3
            while self.emu.clock() < end and self.emu.text_cells() == before:
                self.emu.wait(0.05)
        return n

    # --- battles ----------------------------------------------------------------
    def select(self, item, timeout=10):
        """Move the menu cursor onto `item` (by its text) and confirm."""
        end = self.emu.clock() + timeout
        while self.emu.clock() < end:
            m = self.emu.menu()
            if not m:
                self.emu.wait(0.2)
                continue
            if m['item'] == item:
                self.emu.press('SPACE', gap=0.3)
                return True
            rows = [r for r, text in m['items'] if text == item]
            if not rows:
                raise Blocked(f'menu has no {item!r}: {m["items"]}')
            self.emu.press('UP' if rows[0] < m['row'] else 'DOWN', gap=0.25)
        raise Blocked(f'could not select {item!r}')

    # Command rotation for fights that plain attacks don't end: the Mercenary's fight
    # (D010_X10) only moves on when you cast a spell; Run gives a line and returns.
    FIGHT_PLAN = ['attack'] * 4 + ['magic'] * 2 + ['run']

    def _is_title(self, m):
        names = {x for _, x in m['items']} if m else set()
        return bool(names & {'New Game', 'Watch Opening', '冒険の始めから', 'オープニングから'})

    def fight(self, command=None, max_rounds=60):
        """Fight until the battle ends. Each round uses the next command in FIGHT_PLAN
        (or always `command`, an index into the command menu or its text), takes the
        first choice in any spell/item list, and confirms the target. Battle messages
        are checked like dialogue. Returns rounds fought; raises GameOver if the game
        drops to its title menu."""
        rounds = 0
        seen = set()
        quiet_since = None
        # A battle that shows its panel but never offers a menu or a message would
        # otherwise spin here forever (it blocked whole autoplay runs).
        stuck_budget = 90          # emulated seconds with no menu, message or page
        last_progress = self.emu.clock()
        while rounds < max_rounds:
            if self.emu.clock() - last_progress > stuck_budget:
                raise Blocked(f'battle made no progress for {stuck_budget} s '
                              f'(no menu, no message) after {rounds} rounds')
            if self.check_crash():
                return rounds
            st = self.emu.state()
            if st == 'battle':
                quiet_since = None
                if self.god:
                    self.emu.heal()
                m = self.emu.menu()
                if m and self._is_title(m):
                    self.issue('gameover', 'title menu during a battle')
                    raise GameOver()
                if m and m['items'] and m['col'] >= 70:        # the command menu
                    items = [x for _, x in m['items']]
                    plan = command or self.FIGHT_PLAN[rounds % len(self.FIGHT_PLAN)]
                    pick = {'attack': 0, 'magic': 1, 'run': len(items) - 1}.get(plan)
                    if pick is None:
                        pick = items.index(plan) if plan in items else 0
                    self.select(items[min(pick, len(items) - 1)])
                    self.emu.wait(0.4)
                    for _ in range(3):   # spell/item list, then target: first choice
                        m2 = self.emu.menu()
                        if m2 is None or (m2['items'] and m2['col'] < 70):
                            self.emu.press('SPACE', gap=0.3)
                    rounds += 1
                    last_progress = self.emu.clock()
                    continue
                if m and m['items']:
                    self.emu.press('SPACE', gap=0.3)
                    continue
                # battle messages: anything outside the status/command panels
                d = self.emu.dialogue()
                if d:
                    msg = [l for l, (first, _) in zip(d['lines'], d['cols']) if first < 50]
                    key = tuple(l.strip() for l in msg)
                    if msg and key not in seen:
                        seen.add(key)
                        self.log('  [battle] ' + ' / '.join(key))
                        self.pages.append((self.emu.map_name(), None, msg))
                        for kind, detail in check_page({'lines': msg, 'cols': [(0, 0)] * len(msg),
                                                        'icon_col': None}, self.script)[1]:
                            if kind != 'unmatched':    # battle text is assembled from pieces
                                self.issue(kind, detail)
                    if d['waiting']:
                        self.emu.press('SPACE', gap=0.2)
                    last_progress = self.emu.clock()
                self.emu.wait(0.3)
            elif st == 'dialogue':
                self.read_dialogue(match=False)
                last_progress = self.emu.clock()
            elif st == 'menu':
                m = self.emu.menu()
                if self._is_title(m):
                    self.issue('gameover', 'title menu after a battle')
                    raise GameOver()
                self.emu.press('SPACE', gap=0.3)
            else:
                quiet_since = quiet_since or self.emu.clock()
                if self.emu.clock() - quiet_since > 2.5:
                    self.log(f'  battle over after {rounds} rounds')
                    return rounds
                self.emu.wait(0.3)
        raise Blocked(f'battle still going after {max_rounds} rounds')

    # --- movement that tolerates interruptions ----------------------------------
    def handle_interrupt(self):
        """Deal with whatever took control away from the player. True if anything
        was handled (so a blocked walk is worth retrying)."""
        if self.god:
            # Not only inside fight(): scripted fights and scene battles do not always
            # read as 'battle', and a run died there with every trial annihilated
            # because the heal only ran at the command menu.
            self.emu.heal()
        st = self.emu.state()
        if st == 'battle':
            self.log('  (battle)')
            self.fight()
            return 'battle'
        if st == 'dialogue':
            self.log('  (dialogue opened)')
            self.read_dialogue()
            return 'dialogue'
        return None

    def walk_to(self, x=None, y=None, **kw):
        talks = 0
        for _ in range(8):
            try:
                return self.emu.walk_to(x, y, **kw)
            except Blocked:
                if self.check_crash():
                    raise
                # Give a battle or scene time to show itself: a battle's intro plays
                # for about 3 seconds before its status panel appears.
                what = None
                for _ in range(10):
                    self.emu.wait(0.5)
                    what = self.handle_interrupt()
                    if what or self.check_crash():
                        break
                if what == 'dialogue':
                    talks += 1
                    if talks >= 2:
                        self._sidestep(x, y)
                if what:
                    continue
                raise
        return self.emu.pos()

    def _sidestep(self, x, y):
        """An NPC keeps getting bumped (which starts their conversation again):
        step 20 units sideways, away from the nearest NPC, across the walking
        direction, before trying again."""
        px, py = self.emu.pos()
        near = min(self.emu.objects(), default=None,
                   key=lambda o: abs(o['x'] - px) + abs(o['y'] - py))
        vertical = y is not None and abs(y - py) >= abs((x if x is not None else px) - px)
        if vertical:        # heading up/down: step left or right
            side = 20 if near is None or near['x'] <= px else -20
            self.log(f'  (sidestepping x{side:+d} around an NPC)')
            try:
                self.emu.walk_to(px + side, None)
            except Blocked:
                self.emu.walk_to(px - side, None)
        else:
            side = 20 if near is None or near['y'] <= py else -20
            self.log(f'  (sidestepping y{side:+d} around an NPC)')
            try:
                self.emu.walk_to(None, py + side)
            except Blocked:
                self.emu.walk_to(None, py - side)

    def goto_zone(self, index, kind='step', approach=None, timeout=30, y_first=None):
        """Walk into a trigger zone (see Emu.zones) and wait for whatever it does:
        a map change, a scene, a battle. `approach` = 'UP'/'DOWN'/'LEFT'/'RIGHT'
        enters from that side's opposite (e.g. 'UP' = come from below, moving up);
        by default the zone is entered from the nearer side. With no-clip on the
        straight-line walk ignores walls; without it, walls can block."""
        z = next(z for z in self.emu.zones() if z['kind'] == kind and z['index'] == index)
        cx, cy = (z['x1'] + z['x2']) // 2, (z['y1'] + z['y2']) // 2
        sig0, (px, py) = self.emu.map_sig(), self.emu.pos()
        if approach is None:
            approach = ('DOWN' if py < cy else 'UP') if abs(py - cy) >= abs(px - cx) else \
                       ('RIGHT' if px < cx else 'LEFT')
        # stand 6 units outside the zone on the approach side, lined up with its centre
        if approach in ('UP', 'DOWN'):
            stand = (cx, (z['y2'] + 6) if approach == 'UP' else (z['y1'] - 6))
        else:
            stand = ((z['x2'] + 6) if approach == 'LEFT' else (z['x1'] - 6), cy)
        for attempt in range(6):          # random battles can interrupt; start over
            # y_first: which axis to cover first on the way to the zone (e.g. to pass
            # below an NPC that stands in front of a door)
            yf = (approach in ('LEFT', 'RIGHT')) if y_first is None else y_first
            try:
                self.walk_to(*stand, y_first=yf)
            except Blocked:
                # The spot 6 units outside the zone can be off the map (zones at a map
                # edge) or behind scenery, and then the whole zone is written off as
                # unreachable. Walking onto the zone itself usually still works.
                self.log(f'  (cannot stand off {kind} {index}; going for its centre)')
                self.walk_to(cx, cy, y_first=yf)
            end = self.emu.clock() + timeout
            while self.emu.clock() < end:
                if self.check_crash():
                    return 'crash'
                # battles first: during one the map segment holds battle data
                what = self.handle_interrupt()
                if what == 'battle':
                    break                 # re-approach
                if self.emu.map_sig() != sig0:
                    self.emu.wait(1.0)
                    self.handle_interrupt()
                    return 'map'
                if what == 'dialogue':
                    return 'event'
                before = self.emu.pos()
                self.emu.tap(approach, self.emu.nudge)
                self.emu.wait(0.25)
                if self.emu.pos() == before and self.emu.state() == 'field':
                    # Can't move: often a scene is taking over (the game freezes the
                    # player a moment before its first line appears). Give it 5 s.
                    for _ in range(10):
                        self.emu.wait(0.5)
                        what = self.handle_interrupt()
                        if what or self.emu.map_sig() != sig0:
                            break
                    if what == 'dialogue':
                        return 'event'
                    if what == 'battle':
                        break
                    if self.emu.map_sig() != sig0:
                        return 'map'
                    raise Blocked(f'stuck entering zone {index} at {before}')
            else:
                raise Blocked(f'zone {index} did nothing within {timeout}s')
        raise Blocked(f'zone {index}: interrupted by battles {attempt + 1} times')

    def explore(self, npcs=True, bumps=True, steps=True, per_target_timeout=40):
        """Story-agnostic text coverage for the current map: from a saved state,
        talk to every NPC and walk into every enabled trigger zone one at a time,
        reading and checking whatever text appears, then roll back to the saved
        state before the next target. Returns [(target, outcome, pages, map)].
        Map changes are recorded (and rolled back), so this also maps the exits."""
        self.emu.set_noclip(True)
        base = self.emu.save_state('_explore_base')
        targets = []
        if npcs:
            targets += [('npc', o['slot']) for o in self.emu.objects()]
        for z in self.emu.zones():
            if z['enabled'] and ((z['kind'] == 'bump' and bumps) or (z['kind'] == 'step' and steps)):
                targets.append((z['kind'], z['index']))
        results = []
        start_map = self.emu.map_name()
        for target in targets:
            self.emu.load_state('_explore_base')
            pages_before = len(self.pages)
            self.log(f'explore {start_map}: {target}')
            try:
                if target[0] == 'npc':
                    self.talk_to(target[1])
                    outcome = 'talked'
                else:
                    outcome = self.goto_zone(target[1], kind=target[0], timeout=per_target_timeout)
                    if self.emu.state() == 'dialogue':
                        self.read_dialogue()
            except Blocked as e:
                outcome = f'blocked: {e}'
            except GameOver:
                outcome = 'gameover'
            if self.check_crash():
                outcome = 'CRASH'
            results.append((target, outcome, len(self.pages) - pages_before, self.emu.map_name()))
            self.log(f'  -> {outcome}, {len(self.pages) - pages_before} pages, now on {self.emu.map_name()}')
        self.emu.load_state('_explore_base')
        return results

    def talk_to(self, slot, max_steps=60):
        """Chase a (possibly wandering) NPC, face it and press Space until a
        dialogue opens, then read it. Returns pages read."""
        emu = self.emu
        for _ in range(max_steps):
            o = emu.npc(slot)
            if o is None:
                raise Blocked(f'no object in slot {slot}')
            (mx, my), dx, dy = emu.pos(), o['x'] - emu.pos()[0], o['y'] - emu.pos()[1]
            if abs(dx) <= 5 and abs(dy) <= 7:
                key = (('RIGHT' if dx > 0 else 'LEFT') if abs(dx) > abs(dy)
                       else ('DOWN' if dy > 0 else 'UP'))
                emu.tap(key, emu.nudge)
                emu.press('SPACE', gap=0.3)
                if emu.state() == 'dialogue':
                    return self.read_dialogue()
                continue
            key = (('RIGHT' if dx > 0 else 'LEFT') if abs(dx) > abs(dy) - 3
                   else ('DOWN' if dy > 0 else 'UP'))
            emu.tap(key, min(0.06, max(0.005, (max(abs(dx), abs(dy)) - 4) / 330 * 0.8)))
            self.emu.wait(0.1)
            if emu.state() == 'dialogue':
                return self.read_dialogue()
        raise Blocked(f'could not reach slot {slot}')

    def talk_at(self, key):
        """Face `key` from where we stand and press Space (signs, chests, doors)."""
        self.emu.tap(key, self.emu.nudge)
        self.emu.press('SPACE', gap=0.4)
        if self.emu.state() == 'dialogue':
            return self.read_dialogue()
        return 0

    # --- routes ----------------------------------------------------------------
    def run_step(self, step):
        op = step['op']
        if op == 'walk':
            self.walk_to(step.get('x'), step.get('y'), y_first=step.get('y_first', False))
        elif op == 'exit':
            m = self.emu.exit_via(step['key'])
            if step.get('map') and not same_map(m, step['map']):
                self.issue('route', f'expected map {step["map"]}, got {m}')
        elif op == 'talk':
            self.talk_to(step['slot'])
        elif op == 'talk_at':
            self.talk_at(step['key'])
        elif op == 'read':
            self.read_dialogue()
        elif op == 'press':
            self.emu.press(step['key'], step.get('times', 1), step.get('gap', 0.35))
        elif op == 'wait':
            self.emu.wait(step['seconds'])
        elif op == 'save_state':
            self.emu.save_state(step['name'])
        elif op == 'settle':
            # let scripted scenes play out: read dialogue, fight battles, until the
            # player has had control with nothing happening for `quiet` seconds
            end = self.emu.clock() + step.get('timeout', 300)
            quiet_since = self.emu.clock()
            while self.emu.clock() < end and not self.check_crash():
                if self.handle_interrupt():
                    quiet_since = self.emu.clock()
                elif self.emu.clock() - quiet_since > step.get('quiet', 8):
                    break
                self.emu.wait(0.5)
        elif op == 'explore':
            for target, outcome, pages, where in self.explore():
                self.log(f'  explored {target}: {outcome} ({pages} pages) -> {where}')
        elif op == 'noclip':
            self.emu.set_noclip(step.get('on', True))
        elif op == 'zone':
            r = self.goto_zone(step['index'], step.get('kind', 'step'), step.get('approach'),
                               y_first=step.get('y_first'))
            if step.get('map') and not same_map(self.emu.map_name(), step['map']):
                self.issue('route', f'zone {step["index"]}: expected map {step["map"]}, '
                                    f'got {self.emu.map_name()} ({r})')
        elif op == 'wait_field':
            # cutscenes: keep reading text until the player has control again
            end = self.emu.clock() + step.get('timeout', 60)
            while self.emu.clock() < end and not self.check_crash():
                st = self.emu.state()
                if st == 'dialogue':
                    self.read_dialogue()
                elif self.emu.map_name() and st == 'field':
                    p = self.emu.pos()
                    self.emu.tap('DOWN', 0.01); self.emu.wait(0.3)
                    if self.emu.pos() != p:
                        self.emu.tap('UP', 0.01); self.emu.wait(0.3)
                        break
                else:
                    self.emu.wait(0.5)
        else:
            raise ValueError(f'unknown op {op}')
        if self.emu.state() == 'dialogue':
            self.read_dialogue()

    def start(self, start):
        """Put the game at a route's start point: a named state, or a fresh
        boot to the title menu followed by New Game."""
        if 'state' in start:
            self.emu.load_state(start['state'])
        elif start.get('new_game'):
            self.emu.reset()
            self.emu.wait(3)
            self.emu.title_select(1)

    def play(self, route):
        self.start(route.get('start', {}))
        for i, step in enumerate(route['steps']):
            self.log(f'step {i}: {step}')
            if self.check_crash():
                break
            try:
                self.run_step(step)
            except Blocked as e:
                self.issue('stuck', f'step {i} {step}: {e}')
                break
            except GameOver:
                break
        self.check_crash()

    def report(self, path=None):
        os.makedirs(REPORT_DIR, exist_ok=True)
        path = path or os.path.join(REPORT_DIR, time.strftime('report_%Y%m%d_%H%M%S.txt'))
        with open(path, 'w', encoding='utf-8') as f:
            f.write(f'{len(self.pages)} pages read, {len(self.issues)} issues\n\n')
            for kind, detail, ctx in self.issues:
                f.write(f'[{kind}] {ctx}\n    {detail}\n')
            f.write('\n--- transcript ---\n')
            for m, src, lines in self.pages:
                where = f'{src[0]} {src[1]}' if src else '??'
                f.write(f'{m:6} {where:22} ' + ' / '.join(l.strip() for l in lines) + '\n')
        return path


# --- static checks (no emulator) ----------------------------------------------
BOX_ROWS = 5           # speaker name + 4 lines (confirmed: 02OLB00A 0x3bb's 6th row never shows)


def _cell_terminator(data, off):
    """What follows the text starting at `off` in the original script: 'page' for
    \f (wait for a key, then clear the box), 'line' for \n, else 'other'."""
    i = off
    while i < len(data):
        b = data[i]
        if 0x81 <= b <= 0x9f or b >= 0xe0:
            i += 2
            continue
        if b == 0x5c and i + 1 < len(data):
            if data[i + 1] == 0x66:
                return 'page'
            if data[i + 1] == 0x6e:
                return 'line'
        if b < 0x20 or b == 0x5c:
            return 'other'
        i += 1
    return 'other'


def lint_script(path=os.path.join(HERE, 'KuroNoKen_dump.xlsx')):
    """Find pages that will overflow the text box, without playing. Page structure
    comes from the original script bytes after each cell (\n continues the page,
    \f ends it); each cell is wrapped exactly as reinsert.typeset() will wrap it,
    and pages over BOX_ROWS rows are flagged. Also flags typographic quotes, which
    the game draws as wide full-width glyphs."""
    import openpyxl
    from reinsert import wrap_line
    wb = openpyxl.load_workbook(path, read_only=True)
    cells = {}
    for row in wb['SCNs'].iter_rows(min_row=2, values_only=True):
        if row[0] and row[1]:
            cells.setdefault(row[0], []).append((int(row[1], 16), row[4]))
    out = []
    for fn, rows in cells.items():
        src = os.path.join(HERE, 'original', 'decompressed', fn)
        if not os.path.exists(src):
            src = os.path.join(HERE, 'original', fn)
        if not os.path.exists(src):
            continue
        data = open(src, 'rb').read()
        page, page_start = [], None

        def flush():
            if len(page) > BOX_ROWS:
                out.append(('overfull', fn, f'{page_start:#07x}',
                            f'{len(page)} rows: ' + ' / '.join(l.strip() for l in page)))

        for off, en in rows:
            if page_start is None:
                page_start = off
            if isinstance(en, str) and en.strip() and not en.startswith('='):
                if any(ch in en for ch in '‘’“”'):
                    out.append(('smartquote', fn, f'{off:#07x}', en))
                parts = _SPLIT_RE.split(en)
                for k, part in enumerate(parts):
                    if k:
                        flush()
                        page, page_start = [], off
                    for raw in re.split(r'\[BR\]|\\n', part):
                        page.extend(l.decode('cp932', 'replace') for l in
                                    wrap_line(raw.strip().encode('cp932', 'replace')))
            else:
                page.append('<untranslated>')
            if _cell_terminator(data, off) != 'line':
                flush()
                page, page_start = [], None
        flush()
    return out


# --- recording --------------------------------------------------------------
def record(path, state_name=None, poll=0.03, stop_file=None):
    """Log a human's play as route steps, and check every page of text they read.

    Walk waypoints are taken at each change of movement direction; a map change
    becomes an 'exit' step (and saves a named state, so later replays can start
    there); a dialogue opening becomes 'talk' (nearest NPC), 'talk_at' (facing
    key) or 'read' (it opened by itself). Battles are skipped - replays fight
    them. The route file is rewritten after every step and the text report after
    every page, so nothing is lost if the recorder is killed. Stops on Ctrl+C or
    when `stop_file` (default: <route>.stop) exists."""
    import ctypes
    user32 = ctypes.windll.user32
    tester = Tester()
    emu = tester.emu
    stop_file = stop_file or os.path.splitext(path)[0] + '.stop'
    if os.path.exists(stop_file):
        os.remove(stop_file)
    base = os.path.splitext(os.path.basename(path))[0]
    report_path = os.path.join(REPORT_DIR, base + '_report.txt')
    steps = []
    held = lambda vk: bool(user32.GetAsyncKeyState(vk) & 0x8000)
    dirs = {k: KEYS[k] for k in ('UP', 'DOWN', 'LEFT', 'RIGHT')}
    last_sig, last_pos, last_dir = emu.map_sig(), emu.pos(), None
    last_state, last_key_dir, space_down = emu.state(), None, False
    last_page = None
    start = {'map': emu.map_name(), 'x': last_pos[0], 'y': last_pos[1]}
    if state_name:
        start['state'] = state_name
        emu.save_state(state_name)
    route = {'start': start, 'steps': steps}

    def add(step):
        steps.append(step)
        print('  ', step, flush=True)
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        with open(path, 'w') as f:
            json.dump(route, f, indent=1)

    print(f'recording from {start}; create {stop_file} (or Ctrl+C) to stop', flush=True)
    try:
        while not os.path.exists(stop_file):
            time.sleep(poll)
            st = emu.state()
            pos = emu.pos()
            pressed = [k for k, vk in dirs.items() if held(vk)]
            if pressed:
                last_key_dir = pressed[0]
            if held(KEYS['SPACE']):
                space_down = True
            if st == 'dos':
                tester.check_crash()
                tester.report(report_path)
                print('!! game crashed to DOS', flush=True)
                break

            # passive text checking: every finished page, once
            if st == 'dialogue':
                d = emu.dialogue()
                if d and d['waiting'] and d['lines'] != last_page and any(l.strip() for l in d['lines']):
                    last_page = d['lines']
                    source, problems = check_page(d, tester.script)
                    tester.pages.append((emu.map_name(), source, d['lines']))
                    where = f'{source[0]} {source[1]}' if source else '?'
                    print(f'  [{where}] ' + ' / '.join(l.strip() for l in d['lines']), flush=True)
                    for kind, detail in problems:
                        if kind == 'unmatched' and last_state == 'battle':
                            continue          # battle results are assembled at runtime
                        tester.issue(kind, detail)
                    tester.report(report_path)

            if st == 'battle':
                last_state = st
                continue
            if st == 'field':
                sig = emu.map_sig()
                if sig != last_sig:
                    add({'op': 'exit', 'key': last_key_dir or 'UP', 'map': emu.map_name()})
                    name = f'{base}_{len([s for s in steps if s["op"] == "exit"]):02d}'
                    emu.save_state(name)
                    steps[-1]['state_after'] = name
                    last_sig, last_pos, last_dir = sig, emu.pos(), None
                    last_state = st
                    continue
            if st == 'dialogue' and last_state == 'field':
                near = sorted(emu.objects(), key=lambda o: abs(o['x'] - pos[0]) + abs(o['y'] - pos[1]))
                if last_pos != pos:
                    add({'op': 'walk', 'x': pos[0], 'y': pos[1]})
                if space_down and near and abs(near[0]['x'] - pos[0]) + abs(near[0]['y'] - pos[1]) < 16:
                    add({'op': 'talk', 'slot': near[0]['slot']})
                elif space_down:
                    add({'op': 'talk_at', 'key': last_key_dir})
                else:
                    add({'op': 'read'})
                last_pos = pos
            if st != 'dialogue':
                space_down = space_down and held(KEYS['SPACE'])
            last_state = st
            if st != 'field' or pos == last_pos:
                continue
            dx, dy = pos[0] - last_pos[0], pos[1] - last_pos[1]
            d = 'x' if abs(dx) >= abs(dy) else 'y'
            if last_dir and d != last_dir:
                # direction changed: the previous position is a waypoint
                add({'op': 'walk', 'x': last_pos[0], 'y': last_pos[1], 'y_first': last_dir == 'y'})
            last_dir, last_pos = d, pos
    except KeyboardInterrupt:
        pass
    print(f'wrote {len(steps)} steps to {path}; text report: {tester.report(report_path)}', flush=True)
    print(f'{len(tester.pages)} pages, {len(tester.issues)} issues', flush=True)


if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else 'read'
    if cmd == 'lint':
        found = lint_script()
        for kind, fn, off, detail in found:
            print(f'[{kind}] {fn} {off}: {detail}')
        print(len(found), 'findings')
    elif cmd == 'record':
        name = os.path.splitext(os.path.basename(sys.argv[2]))[0]
        record(sys.argv[2], state_name=name + '_start')
    elif cmd == 'explore':
        t = Tester()
        for row in t.explore():
            print('RESULT', row)
        print('report:', t.report())
        print(f'{len(t.pages)} pages, {len(t.issues)} issues')
    elif cmd in ('read', 'play'):
        t = Tester()
        if cmd == 'read':
            t.read_dialogue()
        else:
            with open(sys.argv[2]) as f:
                t.play(json.load(f))
        print('report:', t.report())
        print(f'{len(t.pages)} pages, {len(t.issues)} issues')
