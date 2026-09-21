"""Merge a workbook exported from the translator's Google Sheet into the repo's dump.

A straight replace loses work in both directions, and breaks reinsertion outright:

* the sheet is ahead on SCNs and BSDs, but the repo is ahead on BD.BIN (2 rows and 33
  translations) and SHINOBU.SMI, so neither file is a superset of the other;
* round-tripping through Google Sheets turns repeated cells into formula references
  (`=E14`), and romtools loads the workbook without `data_only`, so reinsert would
  write the literal text "=E14" into the game for every one of them;
* Google Sheets eats leading whitespace, and the dump's Japanese is padded - 00IPL's
  first line is 36 spaces then the text, and the recorded offset points at the first
  space. 203 rows come back stripped and 23 whitespace-only rows come back empty, which
  moves where the line starts. Running dump_audit on the stripped text does not undo
  this: it "corrects" each offset to where the *stripped* text begins and bakes the
  damage in, so the Japanese has to be put back first.

So: take the sheet's SCNs and BSDs (it has the newer row set there and adds nothing
anywhere else), resolve every English formula to the value it displays, and fill in any
English the repo has and the sheet does not. Rows are aligned per file by their
Japanese, in order, because the offsets differ between the two files until dump_audit
has run. Every other tab is taken from the repo verbatim - the sheet's copies are
strictly older, and BD.BIN's lost rows are easier to keep than to re-insert.

    python tools/merge_dump.py --repo KuroNoKen_dump.xlsx --sheet ~/Downloads/KuroNoKen_dump.xlsx \
                         --out KuroNoKen_dump_merged.xlsx
"""
import argparse
import collections
import difflib
import os
import sys

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # the project root; this file is in tools/
sys.path.insert(0, HERE)


def columns(ws):
    hdr = [c.value for c in list(ws.rows)[0]]
    idx = {n: i for i, n in enumerate(hdr) if isinstance(n, str)}
    return idx.get('Filename'), idx.get('Japanese'), idx.get('English')


def read_rows(ws, fi, ji, ei, sheet_name):
    """[(filename, japanese, english, row_number)] for rows that carry Japanese."""
    out = []
    for row in ws.iter_rows(min_row=2):
        v = [c.value for c in row]
        if ji is None or ji >= len(v):
            continue
        jp = v[ji]
        if jp is None:
            jp = ''
        if not isinstance(jp, str):
            continue
        fn = str(v[fi]) if fi is not None and fi < len(v) and v[fi] else sheet_name
        en = v[ei] if ei is not None and ei < len(v) else None
        out.append((fn, jp, en, row[0].row))
    return out


def has(x):
    return x is not None and str(x).strip() != ''


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--repo', default='KuroNoKen_dump.xlsx')
    ap.add_argument('--sheet', required=True)
    ap.add_argument('--out', default='KuroNoKen_dump_merged.xlsx')
    args = ap.parse_args()

    import openpyxl
    repo = openpyxl.load_workbook(args.repo)
    book = openpyxl.load_workbook(args.sheet)                 # formulas, edited in place
    cached = openpyxl.load_workbook(args.sheet, data_only=True)   # what they display

    FROM_SHEET = ('SCNs', 'BSDs')

    # tabs the translator did not touch: take the repo's, rows and all
    restored = collections.Counter()
    for name in book.sheetnames:
        if name in FROM_SHEET or name not in repo.sheetnames:
            continue
        ws, rs = book[name], repo[name]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)
        for row in rs.iter_rows(min_row=2):
            for cell in row:
                if cell.value is not None:
                    ws.cell(row=cell.row, column=cell.column).value = cell.value
            restored[name] += 1

    resolved = collections.Counter()
    recovered = collections.Counter()
    dropped = collections.Counter()
    repadded = collections.Counter()

    for name in FROM_SHEET:
        ws, wc = book[name], cached[name]
        fi, ji, ei = columns(ws)
        if ei is None:
            continue

        # 1. every English formula becomes the value it displays
        for row in ws.iter_rows(min_row=2):
            if ei >= len(row):
                continue
            cell = row[ei]
            if isinstance(cell.value, str) and cell.value.startswith('='):
                val = wc.cell(row=cell.row, column=ei + 1).value
                if has(val):
                    cell.value = val
                    resolved[name] += 1
                else:
                    # no cached value: Google stores none for a formula showing nothing
                    cell.value = None
                    dropped[name] += 1
            elif cell.value is not None and not isinstance(cell.value, str):
                cell.value = str(cell.value)     # a few cells were typed as numbers

        # 2. fill in English the repo has and this file does not
        if name not in repo.sheetnames:
            continue
        rs = repo[name]
        rfi, rji, rei = columns(rs)
        if rei is None:
            continue
        mine = read_rows(ws, fi, ji, ei, name)
        theirs = read_rows(rs, rfi, rji, rei, name)
        by_file_mine = collections.defaultdict(list)
        by_file_theirs = collections.defaultdict(list)
        for r in mine:
            by_file_mine[r[0]].append(r)
        for r in theirs:
            by_file_theirs[r[0]].append(r)

        for fn, trs in by_file_theirs.items():
            mrs = by_file_mine.get(fn)
            if not mrs:
                continue
            sm = difflib.SequenceMatcher(None, [r[1].strip() for r in trs],
                                         [r[1].strip() for r in mrs], autojunk=False)
            for tag, i1, i2, j1, j2 in sm.get_opcodes():
                if tag != 'equal':
                    continue
                for k in range(i2 - i1):
                    t, m = trs[i1 + k], mrs[j1 + k]
                    if has(t[2]) and not has(ws.cell(row=m[3], column=ei + 1).value):
                        ws.cell(row=m[3], column=ei + 1).value = t[2]
                        recovered[name] += 1
                    if ji is not None and t[1] != m[1] and t[1].strip() == m[1].strip():
                        ws.cell(row=m[3], column=ji + 1).value = t[1]
                        repadded[name] += 1

    book.save(args.out)
    print('tabs taken from the repo verbatim:  %s' % dict(restored))
    print('formulas resolved to their value: %s' % dict(resolved))
    print('formulas showing nothing, cleared: %s' % dict(dropped))
    print('English recovered from the repo:   %s' % dict(recovered))
    print('Japanese whitespace put back:     %s' % dict(repadded))
    print('-> %s' % args.out)


if __name__ == '__main__':
    main()
