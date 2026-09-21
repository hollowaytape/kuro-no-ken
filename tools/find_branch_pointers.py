"""Candidate intra-file pointers that find_pointers' byte patterns miss.

The entry table was one set of them; a block's own branch words (`07 0d 02 00 <addr>`,
`80 00 03 00 00 <addr>`) are another, and they are what still crashes 03YSK01B.SCN.

A word counts as a candidate when it is a valid address in this file, it is not inside
a dumped string, and it points at a **block start** - which is a strong filter, since a
value that merely looks like an address rarely lands exactly on one.

    python tools/find_branch_pointers.py 03YSK01B.SCN [--tuples]
"""
import sys
import os

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, '..'))
from script_map import slot_base, block_starts      # noqa: E402

fn = sys.argv[1]
as_tuples = '--tuples' in sys.argv
data = open(os.path.join(HERE, '..', 'original', 'decompressed', fn), 'rb').read()
base, entries = slot_base(data)
starts = set(block_starts(data, base, entries))

# Byte ranges of the dumped Japanese, so text is never read as a pointer.
import openpyxl                                      # noqa: E402
from rominfo import DUMP_XLS_PATH                    # noqa: E402
spans = []
wb = openpyxl.load_workbook(DUMP_XLS_PATH, read_only=True)
for sheet in wb.sheetnames:
    for row in wb[sheet].iter_rows(min_row=2, values_only=True):
        f, off, jp = row[0], row[1], row[2]
        if f == fn and isinstance(off, str) and isinstance(jp, str):
            a = int(off, 16)
            spans.append((a, a + len(jp.encode('cp932', 'replace'))))


def in_text(i):
    return any(a <= i < b or a <= i + 1 < b for a, b in spans)


known = set()
from rominfo import POINTERS_TO_ADD                  # noqa: E402
for f, loc, _t in POINTERS_TO_ADD:
    if f == fn:
        known.add(loc)

hits = []
for i in range(len(data) - 1):
    if in_text(i) or i in known:
        continue
    w = int.from_bytes(data[i:i + 2], 'little')
    off = w - base
    if off in starts and off > 0:
        hits.append((i, off, data[max(0, i - 5):i].hex(' ')))

print('%s: base %#06x, %d blocks, %d dumped strings' % (fn, base, len(starts), len(spans)))
print('%d candidate pointers (word lands exactly on a block start)' % len(hits))
for i, off, before in hits:
    if as_tuples:
        print("    ('%s', %#0x, %#0x)," % (fn, i, off))
    else:
        print('  %#06x -> %#06x   preceded by %s' % (i, off, before))
