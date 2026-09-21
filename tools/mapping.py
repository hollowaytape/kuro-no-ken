"""Build a "mapping" dump: every line replaced by a FILE-INDEX placeholder.

Two purposes, and the second is the bigger one:

* **a map** - when a placeholder appears on screen you know exactly which dump row is
  being shown, with no text matching involved;
* **a stress test of reinsertion** - 92% of the game's text lives in files that have
  never been reinserted, so their pointers have never had to move. Placeholders make
  every file change length at once, which turns every latent stale pointer into a
  visible one that check_pointers / stale_report / the sweeps can find.

The placeholder keeps the original control codes in place - `\\n`, `\\f`, `\\i2`, colour
codes - because they decide how the box is laid out; only the visible text is replaced.
Each run of text becomes `FILE-N` padded to a little more than the Japanese it replaces,
so files *grow*, which is what real English does to them.

    python tools/mapping.py                    # -> KuroNoKen_dump_mapping.xlsx
    python tools/mapping.py --short            # unpadded ids: files shrink instead
"""
import argparse
import os
import re
import sys

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # the project root; this file is in tools/
sys.path.insert(0, HERE)

CODE = re.compile(r'(\\[a-z](?:\d+(?:,\d+)*)?|\[SPLIT\]|\[[A-Z][A-Za-z0-9]*\])')


def placeholder(jp, tag, pad=True, limit=None):
    """Replace the visible text in `jp`, keeping its control codes where they are."""
    out, n = [], 0
    for part in CODE.split(jp):
        if not part:
            continue
        if CODE.fullmatch(part):
            out.append(part)
            continue
        lead = part[:len(part) - len(part.lstrip())]
        trail = part[len(part.rstrip()):]          # keep it: dropping trailing spaces
        body = part.strip()                        # made every cell a few bytes short,
        if not body:                               # and then everything after it drifted
            out.append(part)
            continue
        n += 1
        ident = '%s-%d' % (tag, n) if n > 1 else tag
        jp_bytes = len(body.encode('cp932', 'replace'))
        if pad:
            # grow the way English does: the Japanese is two bytes a character, so pad
            # to its byte length plus a little. Hub scripts are fixed-length, though
            # (reinsert refuses to make them longer), so there the placeholder has to
            # fit inside what it replaces.
            want = jp_bytes if limit in ('fixed', 'exact') else jp_bytes + 4
            ident = ident + '.' * max(0, want - len(ident))
        if limit in ('fixed', 'exact') and len(ident) > jp_bytes:
            ident = ident[:jp_bytes]
        out.append(lead + ident + trail)
    return ''.join(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default='KuroNoKen_dump_mapping.xlsx')
    ap.add_argument('--short', action='store_true', help='do not pad; files shrink')
    ap.add_argument('--exact', action='store_true',
                    help='pad/truncate to exactly the Japanese byte length: nothing moves, '
                         'so the build needs no pointer relocation at all')
    args = ap.parse_args()

    import openpyxl
    from rominfo import DUMP_XLS_PATH
    from reinsert import FIXED_LENGTH_SCN as FIXED
    wb = openpyxl.load_workbook(DUMP_XLS_PATH)
    counts, filled = {}, 0
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=2):
            fn = row[0].value
            jp = row[2].value if len(row) > 2 else None
            if not fn or not isinstance(jp, str) or not jp:
                continue
            if not jp.strip():
                # Whitespace-only rows (the centred blank lines in the intro) often carry
                # a "[BLANK]" in the English column, which is shorter than the spaces it
                # replaces. For a build meant not to move anything, pass them through.
                if len(row) > 4:
                    row[4].value = jp
                continue
            base = str(fn).rsplit('.', 1)[0]
            counts[base] = counts.get(base, 0) + 1
            tag = '%s-%d' % (base, counts[base])
            if len(row) > 4:
                mode = 'fixed' if fn in FIXED else ('exact' if args.exact else None)
                row[4].value = placeholder(jp, tag, pad=not args.short, limit=mode)
                filled += 1
    wb.save(args.out)
    print('%d placeholder(s) across %d file(s) -> %s' % (filled, len(counts), args.out))
    print('build with:')
    print('   KURO_MAPPING=1 KURO_DUMP_XLS=%s \\' % args.out)
    print('   KURO_POINTER_XLS=KuroNoKen_pointer_dump_mapping.xlsx python reinsert.py')


if __name__ == '__main__':
    main()
