"""Add per-line context columns to the dump workbook, for the translator.

The dump is a flat list of strings in file order, which is enough to reinsert but not
enough to translate well: a cell gives no hint of who is speaking, whether the line is a
speaker label or dialogue, which conversation it belongs to, what has to have happened
for the player to see it, or how much room the text box has left.

All of that is already derivable from the scripts, so this writes it into the workbook
next to each line (columns H onward, left of nothing the reinserter reads):

    Area        the place, from the filename prefix and the Names + Places glossary
    Scene       file + block, so one conversation is one Scene value
    Line        position within the scene ("3 / 12")
    Kind        "name" (a speaker label - keep it short) / "dialogue" / "text"
    Speaker     the label that most recently preceded this line in the same block
    Shown when  what gates this line: "always", "first time only", or a raw flag test
    Sets flags  flags this block sets - i.e. this scene advances the story
    Continues   "-> next cell" when the game does not end the page after this line,
                so the sentence carries on into the following row

Gating comes from the interpreter's three flag instructions (docs/engine_notes.md). What
matters here is which way the *following* code runs, and that was settled by running it,
not by reading the handler: the manor Chancellor's three-page introduction
(03YSK01B 0x241, behind `0c 73 00 <addr>`) plays with flag 0x73 **clear**, and a single
line once it is set. So:

    0c <flag> <addr>   jumps to <addr> if SET   -> code up to <addr> runs when CLEAR
    0d <flag> <addr>   jumps to <addr> if CLEAR -> code up to <addr> runs when SET
    1d <flag>          set the flag

`0c ... 1d <same flag> ... text` is the "first time only" idiom: the lines play once, and
setting the flag skips them from then on.

A print at P is gated by every test at `at` whose jump target t satisfies at < P < t.

    python tools/translator_context.py                       # rewrite KuroNoKen_dump.xlsx in place
    python tools/translator_context.py --out annotated.xlsx  # or to a copy
    python tools/translator_context.py --files 03YSK01A.SCN  # just these, for a quick look

Re-running is idempotent: the columns are recomputed from the scripts, and the
translator's own columns (A-G) are never touched.
"""
import argparse
import collections
import glob
import os
import re
import sys

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # the project root; this file is in tools/
sys.path.insert(0, HERE)

import script_decode as sd                                    # noqa: E402
from script_map import slot_base, block_starts, strings       # noqa: E402

ORIG = os.path.join(HERE, 'original', 'decompressed')

TEST_CLEAR, TEST_SET, SET_FLAG = 0x0c, 0x0d, 0x1d

HEADERS = ['Area', 'Scene', 'Line', 'Kind', 'Speaker', 'Shown when', 'Sets flags',
           'Continues']
FIRST_COL = 8        # column H; A-G are the translator's own


def glossary(wb):
    """-> (area prefix -> name, Japanese character name -> English)."""
    areas, names = {}, {}
    if 'Names + Places' not in wb.sheetnames:
        return areas, names
    for row in wb['Names + Places'].iter_rows(min_row=2, values_only=True):
        src, jp, en, note = (list(row) + [None] * 4)[:4]
        label = en or note
        if src and label:
            areas[str(src).strip().upper()] = str(label).strip()
        elif jp and not src:
            names[str(jp).strip()] = str(en).strip() if en else ''
    return areas, names


def area_for(filename, areas):
    m = re.match(r'[0-9]*[A-Z]+', filename)
    if not m:
        return ''
    key = m.group(0)
    # 10TNII and 26KKRI are sub-areas of 10TNI / 26KKR; fall back to the shorter prefix.
    for k in (key, key[:-1]):
        if k in areas:
            return areas[k]
    return ''


def looks_like_name(jp, names):
    """A speaker label: a bare name on its own line, not a sentence."""
    if not jp:
        return False
    s = jp.strip().strip('　')
    if not s:
        return False
    if s in names:
        return True
    if len(s) > 7:
        return False
    # Dialogue and narration carry quotes or sentence punctuation; a label does not.
    return not any(c in s for c in '「」。、・？！\\')


def flag_gates(insns, base):
    """-> [(at, flag, needs_set, until)] for every flag test, as file offsets.

    `needs_set` is the state the code *after* the test requires - the opposite of the
    state that takes the jump (see the module docstring for how that was confirmed).
    """
    out = []
    for at, op, vals in insns:
        if op in (TEST_CLEAR, TEST_SET) and len(vals) >= 2:
            flag = vals[0][2]
            target = vals[1][2] - base
            out.append((at, flag, op == TEST_SET, target))
    return out


def flags_set(insns):
    return {(at, vals[0][2]) for at, op, vals in insns if op == SET_FLAG and vals}


def page_continues(data, offset):
    """Does the game keep the text box open after the string at `offset`?

    `\\n` (5c 6e) continues the page, `\\f` (5c 66) or any opcode ends it. Same rule as
    reinsert.cell_ends_page, kept here so this tool does not import the reinserter.
    """
    i = offset
    while i < len(data):
        b = data[i]
        if 0x81 <= b <= 0x9f or b >= 0xe0:
            i += 2
            continue
        if b == 0x5c and i + 1 < len(data) and data[i + 1] == 0x6e:
            return True
        if b < 0x20 or b == 0x5c:
            return False
        i += 1
    return False


def context_for(filename, rows, english, areas, names):
    """-> {offset: [Area, Scene, Line, Kind, Speaker, Shown when, Sets flags, Continues]}"""
    path = os.path.join(ORIG, filename)
    if not os.path.exists(path):
        return {}
    data = open(path, 'rb').read()
    base, entries = slot_base(data)
    if base is None:
        return {}
    starts = block_starts(data, base, entries)
    spans = list(zip(starts, starts[1:] + [len(data)]))
    said = strings(data)

    try:
        r = sd.analyse(filename)
        insns = r['insns'] if r and not r.get('guessed_base') else []
    except Exception:
        insns = []
    gates = flag_gates(insns, base)
    sets = flags_set(insns)

    area = area_for(filename, areas)
    by_offset = {}
    for bi, (lo, hi) in enumerate(spans):
        # dump rows this block prints, in file order
        here = sorted(o for o in rows if any(s - 4 <= o <= e for s, e in said if lo <= s < hi))
        if not here:
            continue
        entry = [i for i, a in enumerate(entries) if a == lo]
        scene = '%s #%d%s' % (filename.rsplit('.', 1)[0], bi,
                              ' (entry %d)' % entry[0] if entry else '')
        block_sets = sorted({f for at, f in sets if lo <= at < hi})
        speaker = ''
        for n, off in enumerate(here, 1):
            jp = rows[off]
            is_name = looks_like_name(jp, names)
            if is_name:
                # Prefer a name the translator has already settled on: this row's own
                # English, then the glossary, and only then the raw Japanese.
                bare = jp.strip().strip('　')
                speaker = (english.get(off) or '').strip() or names.get(bare) or bare
            cond = []
            for at, flag, needs_set, until in gates:
                if at < off < until:
                    if not needs_set and flag in block_sets:
                        # gated on F clear by a block that sets F: plays once, then never
                        cond.append('first time only (flag %#x)' % flag)
                    else:
                        cond.append('flag %#x %s' % (flag, 'set' if needs_set else 'clear'))
            by_offset[off] = [
                area,
                scene,
                '%d / %d' % (n, len(here)),
                'name' if is_name else ('dialogue' if jp and '「' in jp else 'text'),
                '' if is_name else speaker,
                '; '.join(sorted(set(cond))) or 'always',
                ', '.join('%#x' % f for f in block_sets),
                '-> next cell' if page_continues(data, off + len(jp.encode('cp932', 'replace'))) else '',
            ]
    return by_offset


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', help='write here instead of in place')
    ap.add_argument('--files', nargs='*', help='only these scripts')
    args = ap.parse_args()

    import openpyxl
    from rominfo import DUMP_XLS_PATH
    wb = openpyxl.load_workbook(DUMP_XLS_PATH)
    areas, names = glossary(wb)
    print('%d areas, %d character names in the glossary' % (len(areas), len(names)))

    ws = wb['SCNs']
    rows_by_file = collections.defaultdict(dict)
    en_by_file = collections.defaultdict(dict)
    cells_by_file = collections.defaultdict(list)
    for i in range(2, ws.max_row + 1):
        fn, off, jp = ws.cell(i, 1).value, ws.cell(i, 2).value, ws.cell(i, 3).value
        en = ws.cell(i, 5).value
        if not fn or off is None:
            continue
        try:
            o = int(str(off), 16)
        except ValueError:
            continue
        rows_by_file[fn][o] = jp if isinstance(jp, str) else ''
        en_by_file[fn][o] = en if isinstance(en, str) else ''
        cells_by_file[fn].append((i, o))

    wanted = set(args.files) if args.files else set(rows_by_file)
    for c, h in enumerate(HEADERS):
        ws.cell(1, FIRST_COL + c).value = h

    filled = files = 0
    for fn in sorted(wanted):
        if fn not in rows_by_file:
            continue
        ctx = context_for(fn, rows_by_file[fn], en_by_file[fn], areas, names)
        if not ctx:
            continue
        files += 1
        for i, o in cells_by_file[fn]:
            vals = ctx.get(o)
            if not vals:
                continue
            for c, v in enumerate(vals):
                ws.cell(i, FIRST_COL + c).value = v
            filled += 1

    for c, width in enumerate((22, 26, 8, 9, 12, 30, 12, 12)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(FIRST_COL + c)].width = width
    out = args.out or DUMP_XLS_PATH
    wb.save(out)
    print('%d rows annotated across %d script(s) -> %s' % (filled, files, out))


if __name__ == '__main__':
    main()
