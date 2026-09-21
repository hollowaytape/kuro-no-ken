"""Every address word in a script that reinsert left behind.

`check_pointers.py` checks the entry table. This checks *every* plausible address in the
file, without knowing which byte patterns hold pointers:

  - align the original and patched files (difflib on bytes) to get an exact
    old-offset -> new-offset map for everything outside the edited strings,
  - for each word in the original that is a valid address into this file and is not
    inside a dumped string, work out where that address should now point,
  - read the same word in the patched file: if it still holds the old value while its
    target moved, reinsert did not know about that pointer.

    python tools/deep_pointer_audit.py [FILE.SCN ...]

With no arguments it audits every translated script. It reports candidates: a value that
merely looks like an address will show up as "stale" even though leaving it alone was
right, so confirm before registering anything (the opcode before it is printed).
"""
import difflib
import glob
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, '..'))
from script_map import slot_base                 # noqa: E402
from check_pointers import as_script             # noqa: E402

ORIG = os.path.join(HERE, '..', 'original', 'decompressed')
PATCHED = os.path.join(HERE, '..', 'patched')
_SPANS = None


def text_spans(fn):
    global _SPANS
    if _SPANS is None:
        import openpyxl
        from rominfo import DUMP_XLS_PATH
        _SPANS = {}
        wb = openpyxl.load_workbook(DUMP_XLS_PATH, read_only=True)
        for sheet in wb.sheetnames:
            for row in wb[sheet].iter_rows(min_row=2, values_only=True):
                f, off, jp = row[0], row[1], row[2]
                if not (f and isinstance(off, str) and isinstance(jp, str)):
                    continue
                try:
                    a = int(off, 16)      # some sheets carry notes in the offset column
                except ValueError:
                    continue
                _SPANS.setdefault(f, []).append((a, a + len(jp.encode('cp932', 'replace'))))
    return _SPANS.get(fn, [])


def offset_map(o, p):
    """old offset -> new offset.

    Equal runs map 1:1, and so does a replacement of the same length - that is what a
    rewritten pointer looks like, and dropping those would blind every caller to
    exactly the bytes reinsert touched.
    """
    m = {}
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, o, p, autojunk=False).get_opcodes():
        if tag == 'equal' or (tag == 'replace' and i2 - i1 == j2 - j1):
            for k in range(i2 - i1):
                m[i1 + k] = j1 + k
    return m


def audit(fn):
    o = open(os.path.join(ORIG, fn), 'rb').read()
    pp = os.path.join(PATCHED, fn)
    if not os.path.exists(pp):
        return []
    p = as_script(open(pp, 'rb').read())
    if p is None or p == o:
        return []
    base, entries = slot_base(o)
    if base is None:
        return []
    spans = text_spans(fn)
    amap = offset_map(o, p)
    stale = []
    for i in range(len(o) - 1):
        if any(a <= i < b or a <= i + 1 < b for a, b in spans):
            continue
        if i not in amap or i + 1 not in amap:
            continue
        w = int.from_bytes(o[i:i + 2], 'little')
        target = w - base
        if not (0x20 <= target < len(o)) or target not in amap:
            continue
        want = base + amap[target]
        j = amap[i]
        if amap[i + 1] != j + 1 or j + 2 > len(p):
            continue
        got = int.from_bytes(p[j:j + 2], 'little')
        if want != w and got == w:                  # target moved, pointer did not
            stale.append((i, target, amap[target], o[max(0, i - 5):i].hex(' ')))
    return stale


def main():
    names = sys.argv[1:] or sorted(os.path.basename(x) for x in glob.glob(os.path.join(ORIG, '*.SCN')))
    total = 0
    for fn in names:
        rows = audit(fn)
        if not rows:
            continue
        total += len(rows)
        print('%s: %d candidate stale pointer(s)' % (fn, len(rows)))
        for i, old, new, before in rows:
            print('   at %#06x: %#06x should now be %#06x   (preceded by %s)'
                  % (i, old, new, before))
    print('%d candidates in %d file(s)' % (total, len(names)))


if __name__ == '__main__':
    main()
